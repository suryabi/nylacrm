"""Tests for new Customer Stock Dashboard logic:
- expected_next_order_date, next_order_urgency, low_stock, priority sort
- Excel export new columns

Assumes QA seed data (QA-CSD-0001..0003) has been loaded via seed_csd_qa.py.
"""
import io
import os
from datetime import date, timedelta

import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
EP = f"{BASE_URL}/api/distributor/customer-stock-dashboard"
LOGIN = f"{BASE_URL}/api/auth/login"


@pytest.fixture(scope="module")
def h():
    r = requests.post(LOGIN, json={"email": "surya.yadavalli@nylaairwater.earth", "password": "test123"})
    assert r.status_code == 200, r.text
    tok = r.json().get("session_token") or r.json().get("token")
    return {"Authorization": f"Bearer {tok}"}


@pytest.fixture(scope="module")
def data(h):
    r = requests.get(EP, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def test_drive_folder_is_first_and_priority(data):
    t0 = data["tiles"][0]
    assert t0["account_name"] == "Drive Folder Test Lead", [t["account_name"] for t in data["tiles"][:5]]
    assert t0["priority"] is True
    assert t0["low_stock"] is True
    assert t0["next_order_urgency"] == "red"
    assert t0["next_order_in_days"] == 5, t0
    assert t0["order_interval_days"] == 30, t0
    assert t0["last_order_quantity"] == 150, t0
    today = date.today()
    assert t0["expected_next_order_date"] == (today + timedelta(days=5)).isoformat()


def test_taj_amber(data):
    taj = next((t for t in data["tiles"] if t["account_name"] == "Taj Sarees"), None)
    assert taj is not None
    assert taj["next_order_urgency"] == "amber", taj
    assert taj["next_order_in_days"] == 9, taj
    assert taj["order_interval_days"] == 30
    assert taj["interval_basis"] == "default (single order)"
    assert taj["low_stock"] is False


def test_no_orders_null(data):
    # Customers with no orders should have nulls
    for t in data["tiles"]:
        if t["last_order_date"] is None:
            assert t["expected_next_order_date"] is None
            assert t["next_order_urgency"] is None
            assert t["next_order_in_days"] is None


def test_last_order_date_iso_10chars(data):
    for t in data["tiles"]:
        if t["last_order_date"]:
            assert len(t["last_order_date"]) == 10 and t["last_order_date"][4] == "-", t


def test_summary_counts(data):
    assert data["summary"]["low_stock"] == 1, data["summary"]
    assert data["summary"]["priority"] == 1


def test_rest_alphabetical_after_priority(data):
    rest = [(t["account_name"] or "").lower() for t in data["tiles"][1:]]
    assert rest == sorted(rest), rest


def test_demo_grand_hotel_green(data):
    gh = next((t for t in data["tiles"] if t["account_name"] == "Demo Grand Hotel (Low Stock)"), None)
    if gh and gh["next_order_urgency"] is not None:
        # Report but do not fail hard on missing seed
        assert gh["next_order_urgency"] == "green", gh


def test_export_has_new_columns(h):
    r = requests.get(f"{EP}/export", headers=h)
    assert r.status_code == 200
    assert "spreadsheet" in r.headers.get("content-type", ""), r.headers
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    # find header row (row with 'Customer' in col A)
    header_row = None
    for row in ws.iter_rows(min_row=1, max_row=10, values_only=False):
        if row[0].value == "Customer":
            header_row = row[0].row
            break
    assert header_row is not None
    headers = [c.value for c in ws[header_row]]
    assert "Expected Next Order" in headers, headers
    assert "Low Stock" in headers, headers
    # find Drive Folder Test Lead row
    dft_row = None
    for row in ws.iter_rows(min_row=header_row + 1, values_only=False):
        if row[0].value == "Drive Folder Test Lead":
            dft_row = row
            break
    assert dft_row is not None, "Drive Folder Test Lead not found in export"
    low_col = headers.index("Low Stock")
    exp_col = headers.index("Expected Next Order")
    assert dft_row[low_col].value == "YES", [c.value for c in dft_row]
    # Expected next order date formatted 'DD Mon YYYY'
    val = dft_row[exp_col].value
    assert isinstance(val, str) and len(val.split()) == 3, val
