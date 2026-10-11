"""Iteration 404 — CIO redesign regression + auth/security essentials."""

import io
import os
import uuid
from typing import Dict, List

import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient


def _read_backend_url() -> str:
    base = (os.environ.get("REACT_APP_BACKEND_URL") or "").strip().strip('"').rstrip("/")
    if base:
        return base
    with open("/app/frontend/.env", "r", encoding="utf-8") as env_file:
        for line in env_file:
            if line.startswith("REACT_APP_BACKEND_URL="):
                return line.split("=", 1)[1].strip().strip('"').rstrip("/")
    raise RuntimeError("REACT_APP_BACKEND_URL missing")


def _read_backend_env(key: str) -> str:
    with open("/app/backend/.env", "r", encoding="utf-8") as env_file:
        for line in env_file:
            if line.startswith(f"{key}="):
                return line.split("=", 1)[1].strip().strip('"').strip("'")
    raise RuntimeError(f"{key} missing in backend/.env")


BASE_URL = _read_backend_url()
TENANT_ID = "nyla-air-water"
ADMIN_EMAIL = "surya.yadavalli@nylaairwater.earth"
ADMIN_PASSWORD = "test123"
MONGO_URL = os.environ.get("MONGO_URL") or _read_backend_env("MONGO_URL")
DB_NAME = os.environ.get("DB_NAME") or _read_backend_env("DB_NAME")

BANGALORE_CODE = "BANG-BAN-A26-001"
EXPECTED_PAYMENT_COUNT = 126
EXPECTED_PAYMENT_GRAND_TOTAL = 816050.20


@pytest.fixture(scope="module")
def mongo_db():
    client = MongoClient(MONGO_URL)
    db = client[DB_NAME]
    yield db
    client.close()


@pytest.fixture(scope="module")
def login_response():
    response = requests.post(
        f"{BASE_URL}/api/auth/login",
        json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD},
        headers={"X-Tenant-ID": TENANT_ID},
        timeout=30,
    )
    if response.status_code != 200:
        pytest.skip(f"Auth failed ({response.status_code}); skipping protected tests")
    return response


@pytest.fixture(scope="module")
def auth_headers(login_response) -> Dict[str, str]:
    payload = login_response.json()
    token = payload.get("token") or payload.get("session_token") or payload.get("access_token")
    if not token:
        pytest.skip("No token in login response")
    return {"Authorization": f"Bearer {token}", "X-Tenant-ID": TENANT_ID}


# Auth baseline checks for this iteration.
def test_auth_cookie_is_http_only(login_response):
    set_cookie = login_response.headers.get("set-cookie", "")
    assert "session_token=" in set_cookie
    assert "HttpOnly" in set_cookie


def test_bcrypt_hash_prefix_is_2b(mongo_db):
    user = mongo_db.users.find_one(
        {"tenant_id": TENANT_ID, "email": ADMIN_EMAIL},
        {"_id": 0, "password": 1},
    )
    assert user and user.get("password")
    assert str(user["password"]).startswith("$2b$")


def test_auth_login_preflight_allows_credentials_for_explicit_origin():
    origin = BASE_URL
    response = requests.options(
        f"{BASE_URL}/api/auth/login",
        headers={
            "Origin": origin,
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type,x-tenant-id",
        },
        timeout=30,
    )
    assert response.status_code in [200, 204]
    assert response.headers.get("access-control-allow-origin") == origin
    assert response.headers.get("access-control-allow-credentials") == "true"


def test_bruteforce_lockout_after_five_failures():
    fake_email = f"iter404-lockout-{uuid.uuid4().hex[:8]}@example.com"
    statuses = []
    for _ in range(8):
        response = requests.post(
            f"{BASE_URL}/api/auth/login",
            json={"email": fake_email, "password": "wrong-pass"},
            headers={"X-Tenant-ID": TENANT_ID},
            timeout=20,
        )
        statuses.append(response.status_code)
    assert statuses[:5] == [401, 401, 401, 401, 401]
    assert 429 in statuses[5:], f"Expected lockout after 5 failures, got {statuses}"


# Customer Invoice Overview API/value contracts after redesign.
def test_cio_payment_totals_and_seed_values_unchanged(auth_headers):
    response = requests.get(
        f"{BASE_URL}/api/reports/customer-invoice-overview",
        headers=auth_headers,
        timeout=60,
    )
    assert response.status_code == 200, response.text
    data = response.json()

    assert len(data.get("months", [])) == 6
    assert data["summary"]["payments"] == EXPECTED_PAYMENT_COUNT
    assert data["payment_grand_total"] == pytest.approx(EXPECTED_PAYMENT_GRAND_TOTAL, abs=0.02)
    assert data["payment_grand_total"] == pytest.approx(sum(data.get("payment_month_totals", [])), abs=0.02)

    row = next((r for r in data["rows"] if r.get("account_code") == BANGALORE_CODE), None)
    assert row is not None, "Bangalore row missing"
    assert len(row.get("payment_months", [])) == 6
    assert row["payment_total"] == pytest.approx(sum(row["payment_months"]), abs=0.02)


def _extract_rows_from_export(content: bytes) -> List[list]:
    ws = load_workbook(io.BytesIO(content)).active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    start = next(i for i, row in enumerate(rows) if row and row[0] == "Customer") + 1
    return [r for r in rows[start:] if r and r[0] is not None]


def test_export_include_payments_contract_unchanged(auth_headers):
    with_payments = requests.get(
        f"{BASE_URL}/api/reports/customer-invoice-overview/export",
        params={"include_payments": "true"},
        headers=auth_headers,
        timeout=90,
    )
    assert with_payments.status_code == 200
    rows_with = _extract_rows_from_export(with_payments.content)
    row_types_with = {r[1] for r in rows_with}
    assert "Invoice" in row_types_with
    assert "Payment received" in row_types_with

    invoice_only = requests.get(
        f"{BASE_URL}/api/reports/customer-invoice-overview/export",
        params={"include_payments": "false"},
        headers=auth_headers,
        timeout=90,
    )
    assert invoice_only.status_code == 200
    rows_invoice_only = _extract_rows_from_export(invoice_only.content)
    non_total_types = {r[1] for r in rows_invoice_only if r[0] != "Total"}
    assert non_total_types == {"Invoice"}


# COO endpoint should remain invoice-only/shared component regression-safe.
def test_coo_endpoint_has_no_payment_fields(auth_headers):
    response = requests.get(
        f"{BASE_URL}/api/reports/customer-order-overview",
        headers=auth_headers,
        timeout=60,
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert len(data.get("months", [])) == 6
    if data.get("rows"):
        sample = data["rows"][0]
        assert "payment_months" not in sample
        assert "payment_total" not in sample