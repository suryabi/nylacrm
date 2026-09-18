"""Test CSD (Customer Stock Dashboard) crate-based forecast + colour logic.

Seeds 4 delivered deliveries for 'Drive Folder Test Lead' with mode-quantity 200 crates,
then hits the /customer-stock-dashboard endpoint and the /export endpoint and asserts the
new crate-based forecast + colour logic.
"""
import os
import io
import asyncio
from datetime import datetime, timezone, timedelta, date

import pytest
import requests
from openpyxl import load_workbook
from motor.motor_asyncio import AsyncIOMotorClient

from dotenv import load_dotenv
load_dotenv("/app/frontend/.env")
BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
MONGO_URL = "mongodb://localhost:27017"
DB_NAME = "test_database"
TENANT = "nyla-air-water"
ACCOUNT_ID = "56b872b1-02e7-4a99-8391-8558c968f450"
CHILD_ID = "ba79db11-532d-49a8-b629-d5a27c04d281"
DISTRIBUTOR_ID = "bb12d90e-4d33-4890-ac5f-17573c551b5c"

CEO_EMAIL = "surya.yadavalli@nylaairwater.earth"
CEO_PASSWORD = "test123"


# ---------------- helpers ----------------
@pytest.fixture(scope="module")
def event_loop():
    loop = asyncio.new_event_loop()
    yield loop
    loop.close()


@pytest.fixture(scope="module")
def db(event_loop):
    client = AsyncIOMotorClient(MONGO_URL)
    return client[DB_NAME]


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"email": CEO_EMAIL, "password": CEO_PASSWORD})
    assert r.status_code == 200, f"login failed {r.status_code} {r.text}"
    body = r.json()
    tok = body.get("session_token") or body.get("token") or body.get("access_token")
    assert tok, f"no token in login response: {body}"
    return tok


@pytest.fixture(scope="module")
def headers(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def seeded(event_loop, db):
    """Seed 4 deliveries and their items. Returns dict with U (units/crate) and parent_sku_id."""
    async def _seed():
        child = await db.child_skus.find_one({"id": CHILD_ID}, {"_id": 0, "parent_sku_id": 1})
        assert child, "child sku missing"
        parent_id = child["parent_sku_id"]
        parent = await db.master_skus.find_one({"id": parent_id}, {"_id": 0, "bottles_per_crate": 1, "packaging_config": 1})
        U = 1
        if parent:
            b = parent.get("bottles_per_crate")
            if not b:
                for pc in (parent.get("packaging_config") or {}).get("production") or []:
                    if pc.get("units_per_package"):
                        b = pc["units_per_package"]
                        break
            if b:
                U = int(b)

        # cleanup any prior
        prior = await db.distributor_deliveries.find({"delivery_number": {"$regex": "^QA-CSD2-"}}, {"_id": 0, "id": 1}).to_list(100)
        prior_ids = [p["id"] for p in prior]
        if prior_ids:
            await db.distributor_delivery_items.delete_many({"delivery_id": {"$in": prior_ids}})
            await db.distributor_deliveries.delete_many({"id": {"$in": prior_ids}})

        # Use 200-crate mode: quantities 200*U, 100*U, 200*U, 200*U
        today = datetime.now(timezone.utc).date()
        offsets = [95, 64, 34, 4]
        crates_list = [200, 100, 200, 200]
        import uuid
        deliveries = []
        items = []
        for i, (off, cr) in enumerate(zip(offsets, crates_list)):
            did = str(uuid.uuid4())
            day = today - timedelta(days=off)
            deliveries.append({
                "id": did, "tenant_id": TENANT, "account_id": ACCOUNT_ID,
                "distributor_id": DISTRIBUTOR_ID,
                "delivery_number": f"QA-CSD2-{i:04d}",
                "delivery_date": day.isoformat(),
                "delivered_at": day.isoformat() + "T10:00:00",
                "status": "delivered",
                "total_quantity": cr * U,
                "distributor_name": "QA Distributor",
            })
            items.append({
                "tenant_id": TENANT, "delivery_id": did,
                "sku_id": parent_id, "quantity": cr * U,
            })
        await db.distributor_deliveries.insert_many(deliveries)
        await db.distributor_delivery_items.insert_many(items)
        return {"U": U, "parent_sku_id": parent_id, "delivery_ids": [d["id"] for d in deliveries]}

    info = event_loop.run_until_complete(_seed())
    yield info
    # teardown
    async def _clean():
        prior = await db.distributor_deliveries.find({"delivery_number": {"$regex": "^QA-CSD2-"}}, {"_id": 0, "id": 1}).to_list(100)
        prior_ids = [p["id"] for p in prior]
        if prior_ids:
            await db.distributor_delivery_items.delete_many({"delivery_id": {"$in": prior_ids}})
            await db.distributor_deliveries.delete_many({"id": {"$in": prior_ids}})
    event_loop.run_until_complete(_clean())


# ---------------- tests ----------------
def _find_tile(tiles, name):
    return next((t for t in tiles if t["account_name"] == name), None)


def test_dashboard_drive_folder_tile(headers, seeded):
    U = seeded["U"]
    r = requests.get(f"{BASE_URL}/api/distributor/customer-stock-dashboard", headers=headers, timeout=60)
    assert r.status_code == 200, r.text
    data = r.json()
    tiles = data["tiles"]
    dft = _find_tile(tiles, "Drive Folder Test Lead")
    assert dft, "Drive Folder Test Lead tile not found"

    # forecast
    assert dft["expected_crates"] == 200, f"expected_crates {dft['expected_crates']}"
    assert "ordered 3 of 4 times" in (dft.get("forecast_basis") or ""), dft.get("forecast_basis")
    # date should be next month on customer's usual day-of-month (gaps 31,30,30 → median 30 > 20)
    assert dft["expected_next_order_date"], "expected_next_order_date should be set"
    exp_date = datetime.fromisoformat(dft["expected_next_order_date"]).date()
    today = datetime.now(timezone.utc).date()
    assert exp_date > today, f"expected date {exp_date} not in future"
    # low stock: 100 units = 100/U crates; U>=1 so 100 <=100 crates < 200 → low stock True
    assert dft["low_stock"] is True
    assert dft["priority"] is True
    assert dft["next_order_urgency"] == "red"
    assert dft["stock_crates"] == round(100 / U, 1), f"stock_crates {dft['stock_crates']} vs {100/U}"
    assert dft["last_order_crates"] == 200
    assert dft["last_order_quantity"] == 200 * U

    # tile MUST be first (priority + A→Z)
    assert tiles[0]["account_id"] == ACCOUNT_ID, f"first tile is {tiles[0]['account_name']}"

    # summary
    assert data["summary"]["low_stock"] >= 1
    assert "stock_crates" in data["summary"]


def test_dashboard_demo_fresh_mart_and_no_branded(headers, seeded):
    r = requests.get(f"{BASE_URL}/api/distributor/customer-stock-dashboard", headers=headers, timeout=60)
    assert r.status_code == 200
    tiles = r.json()["tiles"]

    dfm = _find_tile(tiles, "Demo Fresh Mart (New)")
    if dfm:
        # if it has orders and stock covers expected: green, not priority
        if dfm["expected_next_order_date"]:
            if dfm.get("child_skus"):
                # tracks stock
                assert dfm["next_order_urgency"] in ("green", "red")
                if not dfm["low_stock"]:
                    assert dfm["next_order_urgency"] == "green"
                    assert dfm["priority"] is False

    # A customer with orders but no branded SKU should have urgency 'neutral' and low_stock False
    for t in tiles:
        if t["last_order_date"] and not t["child_skus"]:
            assert t["next_order_urgency"] == "neutral", f"{t['account_name']} urgency"
            assert t["low_stock"] is False
            break

    # customers w/o orders: expected_next_order_date null, urgency null
    for t in tiles:
        if not t["last_order_date"]:
            assert t["expected_next_order_date"] is None
            assert t["next_order_urgency"] is None
            break


def test_dashboard_sort_a_to_z_after_priority(headers, seeded):
    r = requests.get(f"{BASE_URL}/api/distributor/customer-stock-dashboard", headers=headers, timeout=60)
    tiles = r.json()["tiles"]
    priority = [t for t in tiles if t["priority"]]
    non_priority = [t for t in tiles if not t["priority"]]
    assert tiles == priority + non_priority
    non_names = [(t["account_name"] or "").lower() for t in non_priority]
    assert non_names == sorted(non_names)
    prio_names = [(t["account_name"] or "").lower() for t in priority]
    assert prio_names == sorted(prio_names)


def test_export_headers_and_drive_folder_row(headers, seeded):
    r = requests.get(f"{BASE_URL}/api/distributor/customer-stock-dashboard/export", headers=headers, timeout=60)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    # find header row
    header_row_idx = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and row[0] == "Customer":
            header_row_idx = i
            break
    assert header_row_idx is not None
    hdrs = [c for c in ws[header_row_idx] if c.value]
    hdr_values = [c.value for c in ws[header_row_idx]][:12]
    expected = ["Customer", "City", "Category", "Last Order (crates)", "Last Order Date",
                "Expected Next Order", "Expected Crates", "Branded Child SKU",
                "Stock in Hand (crates)", "Stock in Hand (units)", "Low Stock", "Warehouses"]
    assert hdr_values == expected, hdr_values

    # find Drive Folder Test Lead row
    found = False
    for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if row and row[0] == "Drive Folder Test Lead":
            found = True
            assert row[6] == 200, f"Expected Crates cell = {row[6]}"
            assert row[10] == "YES", f"Low Stock cell = {row[10]}"
            break
    assert found, "Drive Folder Test Lead row not found in export"
