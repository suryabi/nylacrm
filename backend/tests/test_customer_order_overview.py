"""Backend regression tests for Customer Order Overview + invoice overview sanity.

Covers report contract, raw-data reconciliation, shared Child SKU grouping,
Excel export, and auth security checks relevant to this feature scope.
"""

import io
import os
import re
import uuid
from datetime import datetime, timezone

import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient


def _read_backend_url() -> str:
    with open("/app/frontend/.env", "r", encoding="utf-8") as f:
        for line in f:
            if line.startswith("REACT_APP_BACKEND_URL="):
                return line.split("=", 1)[1].strip().rstrip("/")
    raise RuntimeError("REACT_APP_BACKEND_URL not found")


BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL") or _read_backend_url()).rstrip("/")
TENANT = "nyla-air-water"
EMAIL = "surya.yadavalli@nylaairwater.earth"
PASSWORD = "test123"

def _read_backend_env(key: str) -> str:
    with open("/app/backend/.env", "r", encoding="utf-8") as f:
        for line in f:
            if line.startswith(f"{key}="):
                return line.split("=", 1)[1].strip().strip('"').strip("'")
    raise RuntimeError(f"{key} not found in /app/backend/.env")


MONGO_URL = os.environ.get("MONGO_URL") or _read_backend_env("MONGO_URL")
DB_NAME = os.environ.get("DB_NAME") or _read_backend_env("DB_NAME")

COMPLETED_STATUSES = {"delivered", "complete", "completed"}


def _month_keys(today, n: int):
    y, m = today.year, today.month
    keys = []
    for _ in range(n):
        keys.append(f"{y:04d}-{m:02d}")
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    return keys[::-1]


def _next_month(month_key: str) -> str:
    year, month = map(int, month_key.split("-"))
    return f"{year + (month == 12):04d}-{1 if month == 12 else month + 1:02d}-01"


def _month_of(delivery: dict) -> str:
    value = delivery.get("delivered_at") or delivery.get("delivery_date") or ""
    return str(value)[:7]


def _units(value) -> int:
    try:
        return max(0, int(value or 0))
    except (TypeError, ValueError):
        return 0


@pytest.fixture(scope="module")
def mongo_db():
    client = MongoClient(MONGO_URL)
    db = client[DB_NAME]
    yield db
    client.close()


@pytest.fixture(scope="module")
def login_response():
    resp = requests.post(
        f"{BASE_URL}/api/auth/login",
        json={"email": EMAIL, "password": PASSWORD},
        headers={"X-Tenant-ID": TENANT},
        timeout=30,
    )
    assert resp.status_code == 200, resp.text
    return resp


@pytest.fixture(scope="module")
def token(login_response):
    payload = login_response.json()
    tok = payload.get("session_token") or payload.get("token") or payload.get("access_token")
    assert tok, payload
    return tok


@pytest.fixture(scope="module")
def headers(token):
    return {"Authorization": f"Bearer {token}", "X-Tenant-ID": TENANT}


def _build_raw_totals(mongo_db: MongoClient):
    """Recompute report totals from eligible raw deliveries + item quantities."""
    now = datetime.now(timezone.utc).date()
    all_keys = _month_keys(now, 7)
    prev_key, keys = all_keys[0], all_keys[1:]
    first_day = f"{prev_key}-01"
    next_day = _next_month(keys[-1])

    accounts = list(
        mongo_db.accounts.find(
            {"tenant_id": TENANT, "status": {"$in": ["active", None]}},
            {"_id": 0, "id": 1},
        )
    )
    account_ids = [a["id"] for a in accounts]

    query = {
        "tenant_id": TENANT,
        "account_id": {"$in": account_ids},
        "status": {"$in": list(COMPLETED_STATUSES)},
        "is_promo": {"$ne": True},
        "$or": [
            {"delivered_at": {"$gte": first_day, "$lt": next_day}},
            {
                "$and": [
                    {
                        "$or": [
                            {"delivered_at": None},
                            {"delivered_at": ""},
                            {"delivered_at": {"$exists": False}},
                        ]
                    },
                    {"delivery_date": {"$gte": first_day, "$lt": next_day}},
                ]
            },
        ],
    }

    deliveries = list(
        mongo_db.distributor_deliveries.find(
            query,
            {
                "_id": 0,
                "id": 1,
                "status": 1,
                "is_promo": 1,
                "delivered_at": 1,
                "delivery_date": 1,
                "total_quantity": 1,
            },
        )
    )
    delivery_ids = [d["id"] for d in deliveries]

    items_map = {}
    if delivery_ids:
        for item in mongo_db.distributor_delivery_items.find(
            {"tenant_id": TENANT, "delivery_id": {"$in": delivery_ids}},
            {"_id": 0, "delivery_id": 1, "quantity": 1},
        ):
            items_map.setdefault(item["delivery_id"], []).append(item)

    month_totals = {k: 0 for k in all_keys}
    delivery_count_in_main_6 = 0
    for d in deliveries:
        m = _month_of(d)
        if m not in all_keys:
            continue
        items = items_map.get(d["id"], [])
        bottles = sum(_units(i.get("quantity")) for i in items)
        if not items:
            bottles = _units(d.get("quantity"))
            if not bottles:
                bottles = _units(d.get("total_quantity"))
        month_totals[m] += bottles
        if m in keys:
            delivery_count_in_main_6 += 1

    return {
        "prev_key": prev_key,
        "keys": keys,
        "month_totals_all": month_totals,
        "delivery_count_in_main_6": delivery_count_in_main_6,
    }


# Auth/security checks relevant for this task's login-dependent APIs.
def test_auth_cookie_set_http_only(login_response):
    set_cookie = login_response.headers.get("set-cookie", "")
    assert "session_token=" in set_cookie
    assert "HttpOnly" in set_cookie


def test_bcrypt_hash_prefix(mongo_db):
    user = mongo_db.users.find_one(
        {"tenant_id": TENANT, "email": EMAIL},
        {"_id": 0, "password": 1},
    )
    assert user and user.get("password")
    assert str(user["password"]).startswith("$2b$")


def test_login_lockout_after_5_failures():
    # Unique email avoids disturbing real users and bypasses prior lock records.
    fake_email = f"lockout-{uuid.uuid4().hex[:8]}@example.com"
    hdr = {"X-Tenant-ID": TENANT}
    for _ in range(5):
        r = requests.post(
            f"{BASE_URL}/api/auth/login",
            json={"email": fake_email, "password": "wrong-pass"},
            headers=hdr,
            timeout=20,
        )
        assert r.status_code == 401

    # Depending on request timing, lock may become visible on the immediate next
    # call or one call later. We require at least one 429 in three attempts.
    statuses = []
    for _ in range(3):
        probe = requests.post(
            f"{BASE_URL}/api/auth/login",
            json={"email": fake_email, "password": "wrong-pass"},
            headers=hdr,
            timeout=20,
        )
        statuses.append(probe.status_code)
        if probe.status_code == 429:
            break
    assert 429 in statuses, f"expected lockout 429 after 5 failures, got {statuses}"


# Customer Order Overview API contract + reconciliation.
def test_customer_order_overview_contract_and_math(headers):
    resp = requests.get(f"{BASE_URL}/api/reports/customer-order-overview", headers=headers, timeout=60)
    assert resp.status_code == 200, resp.text
    data = resp.json()

    assert len(data["months"]) == 6
    keys = [m["key"] for m in data["months"]]
    assert keys == sorted(keys)
    current_key = f"{datetime.now(timezone.utc).year:04d}-{datetime.now(timezone.utc).month:02d}"
    assert keys[-1] == current_key
    assert data["months"][-1]["current"] is True
    assert all(m["current"] is False for m in data["months"][:-1])
    for m in data["months"]:
        assert re.match(r"^[A-Z][a-z]{2} \d{4}$", m["label"]), m["label"]

    rows = data["rows"]
    assert len(rows) > 0
    names = [(r.get("account_name") or "").lower() for r in rows]
    assert names == sorted(names)
    for row in rows:
        assert len(row["months"]) == 6
        assert all(isinstance(v, int) for v in row["months"])
        assert isinstance(row["prev_month"], int)
        assert row["total"] == sum(row["months"])
        assert row["mom_change"] == (row["months"][-1] - row["months"][-2])

    assert len(data["month_totals"]) == 6
    for i in range(6):
        assert data["month_totals"][i] == sum(r["months"][i] for r in rows)
    assert data["prev_month_total"] == sum(r["prev_month"] for r in rows)
    assert data["mom_change_total"] == data["month_totals"][-1] - data["month_totals"][-2]
    assert data["grand_total"] == sum(data["month_totals"])
    assert data["summary"]["rows"] == len(rows)


def test_customer_order_overview_reconciles_with_raw_deliveries(headers, mongo_db):
    api = requests.get(f"{BASE_URL}/api/reports/customer-order-overview", headers=headers, timeout=60)
    assert api.status_code == 200
    report = api.json()

    raw = _build_raw_totals(mongo_db)
    expected_month_totals = [raw["month_totals_all"][k] for k in raw["keys"]]
    assert report["month_totals"] == expected_month_totals
    assert report["prev_month_total"] == raw["month_totals_all"][raw["prev_key"]]
    assert report["summary"]["deliveries"] == raw["delivery_count_in_main_6"]
    assert sum(r.get("delivery_count", 0) for r in report["rows"]) == raw["delivery_count_in_main_6"]


def test_shared_child_sku_grouping_and_dedup(headers):
    report = requests.get(f"{BASE_URL}/api/reports/customer-order-overview", headers=headers, timeout=60).json()
    rows = report["rows"]
    group_rows = [r for r in rows if r.get("row_type") == "brand_group"]
    assert report["summary"]["grouped_rows"] == len(group_rows)
    assert len(group_rows) >= 1

    account_row_ids = {r["account_id"] for r in rows if r.get("row_type") == "account"}
    for group in group_rows:
        members = set(group.get("member_account_ids") or [])
        assert members, "group row must contain member_account_ids"
        assert not (members & account_row_ids), "grouped member leaked as standalone account row"
        assert group.get("member_count") == len(members)
        assert isinstance(group.get("shared_child_sku_names"), list)


def test_customer_order_overview_export_xlsx(headers):
    r = requests.get(f"{BASE_URL}/api/reports/customer-order-overview/export", headers=headers, timeout=90)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers.get("content-type", "")
    assert r.content[:2] == b"PK"

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header = None
    for row in ws.iter_rows(values_only=True):
        if row and row[0] == "Customer / Brand Group":
            header = list(row)
            break
    assert header is not None
    assert header[:7] == [
        "Customer / Brand Group",
        "Row Type",
        "Shared Child SKU(s)",
        "Member Accounts",
        "Code",
        "City",
        "Distributor",
    ]
    assert header[-1] == "Change vs last month"


# Invoice overview regression sanity after shared frontend component refactor.
def test_invoice_overview_and_export_still_work(headers):
    r = requests.get(f"{BASE_URL}/api/reports/customer-invoice-overview", headers=headers, timeout=60)
    assert r.status_code == 200
    data = r.json()
    assert len(data["months"]) == 6
    assert "rows" in data and isinstance(data["rows"], list)

    x = requests.get(f"{BASE_URL}/api/reports/customer-invoice-overview/export", headers=headers, timeout=60)
    assert x.status_code == 200
    assert x.content[:2] == b"PK"
