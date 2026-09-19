"""Backend tests for Brand Group tile on Customer Stock Dashboard.

Validates the new `groups` roll-up returned by
GET /api/distributor/customer-stock-dashboard and the corresponding
`Brand Groups` sheet in the /export xlsx.

Demo dataset (seeded by seed_brand_group_demo.py, MUST remain intact):
  group name "Demo Franchise Group" code DFG
  members  = seed-csd-acct-1, seed-csd-acct-3
  shared child SKU "Demo Franchise – 600 ml" (240 units at Delhi = 20 crates)
"""
import io
import os
import re
import pathlib
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
if not BASE_URL:
    for line in pathlib.Path("/app/frontend/.env").read_text().splitlines():
        m = re.match(r'REACT_APP_BACKEND_URL="?([^"]+)"?', line)
        if m:
            BASE_URL = m.group(1).rstrip("/")
API = f"{BASE_URL}/api"

TENANT = "nyla-air-water"
CEO_EMAIL = "surya.yadavalli@nylaairwater.earth"
CEO_PWD = "test123"

GROUP_NAME = "Demo Franchise Group"
GROUP_CODE = "DFG"
MEMBERS = {"seed-csd-acct-1", "seed-csd-acct-3"}
ACCT_IN_VIEW = "seed-csd-acct-3"
SHARED_CHILD_NAME = "Demo Franchise – 600 ml"


@pytest.fixture(scope="module")
def headers():
    r = requests.post(f"{API}/auth/login", json={"email": CEO_EMAIL, "password": CEO_PWD}, timeout=20)
    assert r.status_code == 200, r.text
    tok = r.json().get("session_token") or r.json().get("token")
    return {"Authorization": f"Bearer {tok}", "X-Tenant-ID": TENANT, "Content-Type": "application/json"}


def _fetch(headers, **params):
    r = requests.get(f"{API}/distributor/customer-stock-dashboard", headers=headers, params=params, timeout=30)
    assert r.status_code == 200, r.text
    return r.json()


def _demo_group(data):
    g = next((g for g in data.get("groups") or [] if g.get("name") == GROUP_NAME), None)
    assert g, f"Demo Franchise Group missing in groups={[x.get('name') for x in data.get('groups') or []]}"
    return g


# --- 1. dashboard group fields --------------------------------------------------
def test_demo_group_shape_and_values(headers):
    data = _fetch(headers)
    g = _demo_group(data)
    assert g.get("group_id")
    assert g["code"] == GROUP_CODE
    assert g["total_members"] == 2
    assert len(g["members"]) == 2
    assert {m["account_id"] for m in g["members"]} == MEMBERS
    for m in g["members"]:
        for k in ("account_id", "account_name", "last_order_crates", "last_order_date",
                  "expected_crates", "expected_next_order_date", "low_stock", "own_skus"):
            assert k in m, f"member missing key {k}: {m}"

    assert len(g["shared_skus"]) == 1
    s = g["shared_skus"][0]
    assert s["shared"] is True
    assert s["stock_crates"] == 20
    assert s["stock_in_hand"] == 240
    assert s["child_sku_name"] == SHARED_CHILD_NAME

    assert g["stock_crates"] == 20
    # sum & max member roll-ups
    lc = [m["last_order_crates"] for m in g["members"] if m["last_order_crates"] is not None]
    ld = [m["last_order_date"] for m in g["members"] if m["last_order_date"]]
    assert g["last_orders_crates"] == round(sum(lc), 1)
    assert g["last_order_date"] == max(ld)
    ec = [m["expected_crates"] for m in g["members"] if m["expected_crates"] is not None]
    assert g["expected_crates"] == (round(sum(ec), 1) if ec else None) or g["expected_crates"] == int(round(sum(ec), 1))
    ed = [m["expected_next_order_date"] for m in g["members"] if m["expected_next_order_date"]]
    if ed:
        assert g["expected_next_order_date"] == min(ed)
    # history totals: 4 orders * <=6 sampled per member; the seeded expectation is 50 crates / 4 orders
    assert g["history_orders"] >= len(lc)
    assert g["history_crates"] >= 0
    # low_stock true when shared stock < expected
    if g["expected_crates"] is not None:
        assert g["low_stock"] is bool(g["stock_crates"] < g["expected_crates"])
    assert g["priority"] == g["low_stock"]
    if g["low_stock"]:
        assert g["next_order_urgency"] in ("red",)
    elif g["shared_skus"]:
        assert g["next_order_urgency"] in ("green", "red")


def test_summary_and_tile_brand_group_ids(headers):
    data = _fetch(headers)
    sm = data["summary"]
    assert sm.get("groups") == len(data["groups"])
    assert sm.get("grouped_customers") == sum(len(g["members"]) for g in data["groups"])
    # member tiles carry brand_group_ids; non-members are []
    gid = _demo_group(data)["group_id"]
    by_id = {t["account_id"]: t for t in data["tiles"]}
    for mid in MEMBERS:
        assert mid in by_id, f"member tile {mid} missing from `tiles`"
        assert gid in by_id[mid].get("brand_group_ids") or []
    non_member = next((t for t in data["tiles"] if t["account_id"] not in MEMBERS), None)
    if non_member:
        assert non_member.get("brand_group_ids") == []


def test_shared_stock_counted_once_in_summary(headers):
    """summary.stock_crates must NOT double the shared SKU across members."""
    data = _fetch(headers)
    child_ids_seen = set()
    expected_units = 0
    expected_crates = 0.0
    for t in data["tiles"]:
        for s in t["child_skus"]:
            if s["child_sku_id"] in child_ids_seen:
                continue
            child_ids_seen.add(s["child_sku_id"])
            expected_units += s["stock_in_hand"]
            expected_crates += s["stock_crates"]
    sm = data["summary"]
    assert sm["stock_in_hand"] == expected_units
    assert abs(sm["stock_crates"] - round(expected_crates, 1)) < 0.05


# --- 2. filters -----------------------------------------------------------------
def test_filter_by_single_member_account(headers):
    data = _fetch(headers, account_id=ACCT_IN_VIEW)
    g = _demo_group(data)
    assert g["total_members"] == 2
    assert len(g["members"]) == 1
    assert g["members"][0]["account_id"] == ACCT_IN_VIEW


def test_filter_distributor_not_serving_hides_group(headers):
    # find some distributor not serving either member (factory/non-Delhi)
    r = requests.get(f"{API}/distributors", headers=headers, timeout=20)
    if r.status_code != 200:
        pytest.skip(f"distributors list unavailable: {r.status_code}")
    for d in r.json().get("distributors") or r.json() or []:
        did = d.get("id")
        if not did:
            continue
        data = _fetch(headers, distributor_id=did)
        member_tiles = [t for t in data["tiles"] if t["account_id"] in MEMBERS]
        if not member_tiles:
            assert not any(g.get("name") == GROUP_NAME for g in data["groups"]), \
                "group returned though no members are in view"
            return
    pytest.skip("every distributor serves at least one Demo member; cannot test empty-member case")


# --- 3. new group with no shared SKU -------------------------------------------
def test_group_without_shared_sku(headers):
    payload = {"name": "QA Empty Group Tile", "account_ids": list(MEMBERS)}
    r = requests.post(f"{API}/brand-groups", headers=headers, json=payload, timeout=20)
    assert r.status_code in (200, 201), r.text
    gid = r.json()["id"]
    try:
        data = _fetch(headers)
        g = next((g for g in data["groups"] if g["group_id"] == gid), None)
        assert g, "new empty group not returned"
        assert g["shared_skus"] == []
        assert g["stock_crates"] == 0
        assert g["low_stock"] is False
        assert g["next_order_urgency"] in (None, "neutral")
    finally:
        d = requests.delete(f"{API}/brand-groups/{gid}", headers=headers, timeout=20)
        assert d.status_code in (200, 204), d.text


# --- 4. export xlsx -------------------------------------------------------------
def test_export_has_brand_groups_sheet(headers):
    r = requests.get(f"{API}/distributor/customer-stock-dashboard/export", headers=headers, timeout=60)
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers.get("content-type", "")
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    assert "Brand Groups" in wb.sheetnames, f"sheets={wb.sheetnames}"
    ws = wb["Brand Groups"]
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    # locate header row containing "Brand Group"
    hdr_idx = next((i for i, r in enumerate(rows) if r and r[0] == "Brand Group"), None)
    assert hdr_idx is not None, "header row not found"
    # find Demo Franchise Group data row
    demo_row = next((r for r in rows[hdr_idx + 1:] if r and r[0] == GROUP_NAME), None)
    assert demo_row, "Demo Franchise Group row missing"
    assert demo_row[1] == "2 of 2"
    assert demo_row[2] == SHARED_CHILD_NAME
    assert demo_row[3] == 20
    assert demo_row[4] == 240
    # member arrow rows exist below
    member_rows = [r for r in rows[hdr_idx + 1:] if r and isinstance(r[1], str) and r[1].startswith("↳ ")]
    assert len(member_rows) >= 2, f"expected member '↳' rows, got {member_rows}"
