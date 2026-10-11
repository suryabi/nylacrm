"""Iteration 403 — Customer Invoice Overview payment-row backend contracts."""
import io
import os
from typing import Dict, List

import pytest
import requests
from openpyxl import load_workbook


def _read_backend_url() -> str:
    base = (os.environ.get("REACT_APP_BACKEND_URL") or "").strip().strip('"').rstrip("/")
    if base:
        return base
    with open("/app/frontend/.env", "r", encoding="utf-8") as env_file:
        for line in env_file:
            if line.startswith("REACT_APP_BACKEND_URL="):
                return line.split("=", 1)[1].strip().strip('"').rstrip("/")
    raise RuntimeError("REACT_APP_BACKEND_URL missing")


BASE_URL = _read_backend_url()
TENANT_ID = "nyla-air-water"
ADMIN_EMAIL = "surya.yadavalli@nylaairwater.earth"
ADMIN_PASSWORD = "test123"
BANGALORE_UUID = "91e3dd51-5d07-4f75-a702-02994b0e770f"
BANGALORE_CODE = "BANG-BAN-A26-001"
EXPECTED_PAYMENTS = [1111.0, 0.0, 0.0, 0.0, 5432.0, 12345.5]
EXPECTED_PREV = 500.0
EXPECTED_TOTAL = 18888.5
EXPECTED_MOM = 6913.5
EXPECTED_COUNT = 4


@pytest.fixture(scope="module")
def auth_headers() -> Dict[str, str]:
    login = requests.post(
        f"{BASE_URL}/api/auth/login",
        json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD},
        headers={"X-Tenant-ID": TENANT_ID},
        timeout=30,
    )
    assert login.status_code == 200, login.text
    payload = login.json()
    token = payload.get("token") or payload.get("session_token") or payload.get("access_token")
    assert token, payload
    return {"Authorization": f"Bearer {token}", "X-Tenant-ID": TENANT_ID}


@pytest.fixture(scope="module")
def cio_data(auth_headers) -> dict:
    response = requests.get(
        f"{BASE_URL}/api/reports/customer-invoice-overview",
        headers=auth_headers,
        timeout=60,
    )
    assert response.status_code == 200, response.text
    data = response.json()
    assert len(data.get("months", [])) == 6
    return data


# Core response contract + aggregate consistency for invoice/payment dual metrics
def test_cio_row_contract_and_global_reconciliation(cio_data):
    rows = cio_data["rows"]
    assert rows, "Expected non-empty report rows"

    for row in rows:
        assert len(row.get("months", [])) == 6
        assert len(row.get("payment_months", [])) == 6
        assert isinstance(row.get("payment_prev_month"), (int, float))
        assert isinstance(row.get("payment_total"), (int, float))
        assert isinstance(row.get("payment_mom_change"), (int, float))
        assert isinstance(row.get("payment_count"), int)
        assert row["total"] == pytest.approx(round(sum(row["months"]), 2), abs=0.02)
        assert row["payment_total"] == pytest.approx(round(sum(row["payment_months"]), 2), abs=0.02)
        assert row["mom_change"] == pytest.approx(round(row["months"][-1] - row["months"][-2], 2), abs=0.02)
        assert row["payment_mom_change"] == pytest.approx(
            round(row["payment_months"][-1] - row["payment_months"][-2], 2), abs=0.02
        )

    month_totals = [round(sum(r["months"][i] for r in rows), 2) for i in range(6)]
    payment_month_totals = [round(sum(r["payment_months"][i] for r in rows), 2) for i in range(6)]

    assert cio_data["month_totals"] == pytest.approx(month_totals, abs=0.02)
    assert cio_data["payment_month_totals"] == pytest.approx(payment_month_totals, abs=0.02)
    assert cio_data["prev_month_total"] == pytest.approx(round(sum(r["prev_month"] for r in rows), 2), abs=0.02)
    assert cio_data["payment_prev_month_total"] == pytest.approx(
        round(sum(r["payment_prev_month"] for r in rows), 2), abs=0.02
    )
    assert cio_data["grand_total"] == pytest.approx(round(sum(cio_data["month_totals"]), 2), abs=0.02)
    assert cio_data["payment_grand_total"] == pytest.approx(round(sum(cio_data["payment_month_totals"]), 2), abs=0.02)

    summary = cio_data["summary"]
    assert summary["paid_customers"] == sum(1 for row in rows if row["payment_total"] > 0)
    assert summary["payments"] == sum(row["payment_count"] for row in rows)


# Seed fixture checks for payment inclusion/exclusion behavior and month bucketing
def test_bangalore_seed_fixture_values(cio_data):
    row = next(
        (
            r
            for r in cio_data["rows"]
            if r.get("account_id") == BANGALORE_UUID or r.get("account_code") == BANGALORE_CODE
        ),
        None,
    )
    assert row is not None, "Bangalore Tech Park row not found"
    assert row["payment_months"] == pytest.approx(EXPECTED_PAYMENTS, abs=0.01)
    assert row["payment_prev_month"] == pytest.approx(EXPECTED_PREV, abs=0.01)
    assert row["payment_total"] == pytest.approx(EXPECTED_TOTAL, abs=0.01)
    assert row["payment_mom_change"] == pytest.approx(EXPECTED_MOM, abs=0.01)
    assert row["payment_count"] == EXPECTED_COUNT


def _fetch_export(auth_headers, include_payments: bool):
    response = requests.get(
        f"{BASE_URL}/api/reports/customer-invoice-overview/export",
        params={"include_payments": str(include_payments).lower()},
        headers=auth_headers,
        timeout=90,
    )
    assert response.status_code == 200, response.text
    assert "spreadsheetml" in response.headers.get("content-type", "")
    assert response.content[:2] == b"PK"
    return load_workbook(io.BytesIO(response.content)).active


def _sheet_rows(ws) -> List[list]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _extract_table_rows(rows: List[list]) -> List[list]:
    start = next(i for i, row in enumerate(rows) if row and row[0] == "Customer") + 1
    return [r for r in rows[start:] if r and r[0] is not None]


# Excel contract for include_payments true/false and row pairing/order integrity
def test_export_include_payments_true_false(auth_headers, cio_data):
    ws_with_payments = _fetch_export(auth_headers, include_payments=True)
    ws_invoice_only = _fetch_export(auth_headers, include_payments=False)

    rows_with_payments = _extract_table_rows(_sheet_rows(ws_with_payments))
    rows_invoice_only = _extract_table_rows(_sheet_rows(ws_invoice_only))

    account_count = len(cio_data["rows"])
    invoice_rows = [r for r in rows_with_payments if r[1] == "Invoice" and r[0] != "Total"]
    payment_rows = [r for r in rows_with_payments if r[1] == "Payment received" and r[0] != "Total"]
    assert len(invoice_rows) == account_count
    assert len(payment_rows) == account_count

    # Pairing: every customer has Invoice immediately followed by Payment received
    for index in range(0, account_count * 2, 2):
        first = rows_with_payments[index]
        second = rows_with_payments[index + 1]
        assert first[1] == "Invoice"
        assert second[1] == "Payment received"
        assert first[0] == second[0]

    total_rows = [r for r in rows_with_payments if r[0] == "Total"]
    assert len(total_rows) == 2
    assert total_rows[0][1] == "Invoice"
    assert total_rows[1][1] == "Payment received"
    assert total_rows[0][5:11] == pytest.approx(cio_data["month_totals"], abs=0.02)
    assert total_rows[1][5:11] == pytest.approx(cio_data["payment_month_totals"], abs=0.02)

    bangalore_payment = next(
        (r for r in payment_rows if r[0] and BANGALORE_CODE in str(r[2] or "")),
        None,
    )
    assert bangalore_payment is not None, "Bangalore payment row missing in include_payments=true export"
    assert bangalore_payment[5:11] == pytest.approx(EXPECTED_PAYMENTS, abs=0.01)
    assert float(bangalore_payment[11]) == pytest.approx(EXPECTED_MOM, abs=0.01)

    invoice_only_types = {r[1] for r in rows_invoice_only if r[0] != "Total"}
    assert invoice_only_types == {"Invoice"}
    invoice_only_totals = [r for r in rows_invoice_only if r[0] == "Total"]
    assert len(invoice_only_totals) == 1
    assert invoice_only_totals[0][1] == "Invoice"
