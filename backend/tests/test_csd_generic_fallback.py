"""Tests for Customer Stock Dashboard 'generic fallback' behavior.

Verifies:
- stock_kind in {'branded','generic',None} per tile
- Generic tiles for Empire Restaurant (Brian distributor scope), Oceans Fresh Seafood (all scope), Demo Grand Hotel (Low Stock)
- child_sku_id format 'parent:<sku>:<distributor_id|all>'
- Generic stock math matches sum over distributor_stock (child_sku_id null/'') for that distributor
- summary.generic_customers == 7, summary.branded_skus == 17
- Export header contains 'SKU (branded child / generic parent)' and (generic) suffix, 'No SKU' for empty
"""
import io
import os
import re
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
TENANT = "nyla-air-water"
CREDS = {"email": "surya.yadavalli@nylaairwater.earth", "password": "test123"}


@pytest.fixture(scope="module")
def client():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json", "X-Tenant-ID": TENANT})
    # Try login endpoints
    r = s.post(f"{BASE_URL}/api/auth/login", json=CREDS)
    if r.status_code != 200:
        pytest.skip(f"Login failed: {r.status_code} {r.text[:200]}")
    tok = r.json().get("session_token") or r.json().get("token") or r.json().get("access_token")
    if tok:
        s.headers["Authorization"] = f"Bearer {tok}"
    return s


@pytest.fixture(scope="module")
def dashboard(client):
    r = client.get(f"{BASE_URL}/api/distributor/customer-stock-dashboard")
    assert r.status_code == 200, r.text[:500]
    return r.json()


def _tile(dashboard, name_frag):
    hits = [t for t in dashboard["tiles"] if name_frag.lower() in (t.get("account_name") or "").lower()]
    assert hits, f"No tile matching {name_frag!r}"
    return hits[0]


# ----- Structural checks -----

def test_stock_kind_domain(dashboard):
    for t in dashboard["tiles"]:
        assert t.get("stock_kind") in ("branded", "generic", None), (t["account_name"], t.get("stock_kind"))


def test_branded_tiles_unchanged(dashboard):
    branded = [t for t in dashboard["tiles"] if t["stock_kind"] == "branded"]
    assert branded, "Expected branded tiles"
    for t in branded:
        for s in t["child_skus"]:
            assert not s.get("generic"), f"{t['account_name']} branded row has generic=True"
            # child_sku_id must not be a synthetic parent key
            assert not str(s["child_sku_id"]).startswith("parent:"), s


def test_empty_sku_tiles(dashboard):
    empty = [t for t in dashboard["tiles"] if t["stock_kind"] is None]
    for t in empty:
        assert t["child_skus"] == []


# ----- Empire (Brian distributor scope) -----

def test_empire_generic_two_rows(dashboard):
    t = _tile(dashboard, "Empire Restaurant")
    assert t["stock_kind"] == "generic", t
    names = sorted(s["child_sku_name"] for s in t["child_skus"])
    assert len(t["child_skus"]) == 2, names
    # SKU names expected
    joined = " | ".join(names).lower()
    assert "330" in joined and "600" in joined, names
    # Every row: generic=True, scope='distributor', synthetic id, wh only from that distributor's locations
    for s in t["child_skus"]:
        assert s["generic"] is True
        assert s["scope"] == "distributor"
        assert re.match(r"^parent:[^:]+:[^:]+$", s["child_sku_id"]), s["child_sku_id"]
        assert s["child_sku_id"].endswith(f":{t['distributor_id']}")
    # crates roughly 5.7 + 118.5 = 124.2
    total_crates = round(sum(s["stock_crates"] for s in t["child_skus"]), 1)
    assert 120.0 <= total_crates <= 130.0, total_crates


def test_oceans_generic_all_warehouses(dashboard):
    t = _tile(dashboard, "Oceans Fresh Seafood")
    assert t["stock_kind"] == "generic"
    assert t["distributor_name"] in ("Unassigned", None) or t["distributor_id"] is None
    for s in t["child_skus"]:
        assert s["scope"] == "all"
        assert s["child_sku_id"].endswith(":all")


def test_demo_grand_generic_low_stock(dashboard):
    t = _tile(dashboard, "Demo Grand Hotel")
    assert t["stock_kind"] == "generic", t
    assert t["low_stock"] is True, t


# ----- Summary -----

def test_summary_counts(dashboard):
    s = dashboard["summary"]
    assert s["generic_customers"] == 7, s["generic_customers"]
    branded_sku_count = sum(len(t["child_skus"]) for t in dashboard["tiles"] if t["stock_kind"] == "branded")
    assert s["branded_skus"] == branded_sku_count, (s["branded_skus"], branded_sku_count)
    assert s["branded_skus"] == 17, s["branded_skus"]


def test_summary_stock_crates_dedupes(dashboard):
    # Recompute the same way backend does: unique by child_sku_id
    uniq = {}
    for t in dashboard["tiles"]:
        for row in t["child_skus"]:
            uniq[row["child_sku_id"]] = row
    expected = round(sum(v["stock_crates"] for v in uniq.values()), 1)
    assert dashboard["summary"]["stock_crates"] == expected


# ----- Generic stock math cross-check against Mongo -----

def test_generic_math_against_mongo(dashboard):
    import sys, asyncio
    sys.path.insert(0, "/app/backend")
    from database import db  # noqa
    t = _tile(dashboard, "Empire Restaurant")
    dist_id = t["distributor_id"]
    assert dist_id
    row = t["child_skus"][0]
    sid = row["sku_id"]

    async def _sum():
        s = 0
        async for r in db.distributor_stock.find(
            {"tenant_id": TENANT, "sku_id": sid, "child_sku_id": {"$in": [None, ""]},
             "quantity": {"$gt": 0}, "distributor_id": dist_id},
            {"_id": 0, "quantity": 1},
        ):
            s += int(r.get("quantity") or 0)
        return s

    total = asyncio.get_event_loop().run_until_complete(_sum())
    assert row["stock_in_hand"] == total, (row["stock_in_hand"], total)
    upc = row["units_per_crate"] or 1
    assert row["stock_crates"] == round(total / upc, 1)


# ----- Export -----

def test_export_header_and_suffix(client, dashboard):
    r = client.get(f"{BASE_URL}/api/distributor/customer-stock-dashboard/export")
    assert r.status_code == 200, r.text[:300]
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    # Locate header row (row with the SKU column)
    header_row = None
    sku_col = None
    for row in ws.iter_rows(values_only=False):
        for cell in row:
            if cell.value == "SKU (branded child / generic parent)":
                header_row = cell.row
                sku_col = cell.column
                break
        if header_row:
            break
    assert header_row, "SKU header not found"
    # Scan rows for '(generic)' suffix and 'No SKU'
    values = [ws.cell(row=r, column=sku_col).value for r in range(header_row + 1, ws.max_row + 1)]
    joined = " | ".join(str(v) for v in values if v)
    assert "(generic)" in joined, "Missing '(generic)' suffix in export SKU column"
    assert "No SKU" in joined, "Missing 'No SKU' for empty SKU customers"
