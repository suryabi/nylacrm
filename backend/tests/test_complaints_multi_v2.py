"""Regression tests for Issues (customer_complaints) module - multi-user/multi-department,
close/reopen, no category, import/export XLSX. Iteration 340."""
import io
import os
import pytest
import requests
from openpyxl import Workbook, load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://design-workflow-hub-11.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

EMAIL = "surya.yadavalli@nylaairwater.earth"
PASSWORD = "test123"


@pytest.fixture(scope="module")
def session():
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": EMAIL, "password": PASSWORD}, timeout=30)
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text}"
    tok = r.json().get("session_token") or r.json().get("token")
    assert tok, f"no session_token in {r.json()}"
    s.headers.update({"Authorization": f"Bearer {tok}", "Content-Type": "application/json"})
    return s


@pytest.fixture(scope="module")
def options(session):
    r = session.get(f"{API}/complaints/meta/options")
    assert r.status_code == 200
    return r.json()


def _find_two_users(options):
    users = options.get("users", [])
    assert len(users) >= 2, "need >=2 users to test multi-assign"
    return users[0], users[1]


# ────────────────────── META / CREATE / UPDATE ──────────────────────

def test_meta_options_no_categories(options):
    assert "categories" not in options
    assert "departments" in options
    assert "users" in options


def test_create_multi_dep_multi_assignee(session, options):
    u1, u2 = _find_two_users(options)
    payload = {
        "title": "TEST_multi_v2_create",
        "priority": "high",
        "departments": ["Production", "Quality"],
        "assignees": [{"id": u1["id"], "name": u1["name"]}, {"id": u2["id"], "name": u2["name"]}],
    }
    r = session.post(f"{API}/complaints", json=payload)
    assert r.status_code == 200, r.text
    d = r.json()
    assert "category" not in d or d.get("category") in (None, "")
    assert d["departments"] == ["Production", "Quality"]
    assert d["department"] == "Production"  # legacy mirror
    assert len(d["assignees"]) == 2
    assert d["assigned_to"] == u1["id"]
    assert d["assigned_to_name"] == u1["name"]
    cid = d["id"]

    # PUT update departments
    r = session.put(f"{API}/complaints/{cid}", json={"departments": ["Quality", "Sales"]})
    assert r.status_code == 200
    d = r.json()
    assert d["departments"] == ["Quality", "Sales"]
    assert d["department"] == "Quality"

    # PUT update assignees to one only
    r = session.put(f"{API}/complaints/{cid}", json={"assignees": [{"id": u2["id"], "name": u2["name"]}]})
    assert r.status_code == 200
    d = r.json()
    assert len(d["assignees"]) == 1
    assert d["assigned_to"] == u2["id"]

    # Close
    r = session.put(f"{API}/complaints/{cid}", json={"status": "closed"})
    assert r.status_code == 200
    d = r.json()
    assert d["status"] == "closed"
    assert d.get("closed_at")

    # Reopen (open from closed)
    r = session.put(f"{API}/complaints/{cid}", json={"status": "open"})
    assert r.status_code == 200
    d = r.json()
    assert d["status"] == "open"
    assert d.get("reopened_at")

    # Filter by department
    r = session.get(f"{API}/complaints", params={"department": "Sales"})
    assert r.status_code == 200
    ids = [x["id"] for x in r.json()["data"]]
    assert cid in ids

    # Filter by assigned_to id
    r = session.get(f"{API}/complaints", params={"assigned_to": u2["id"]})
    assert r.status_code == 200
    ids = [x["id"] for x in r.json()["data"]]
    assert cid in ids

    # unassigned - remove all assignees then check
    session.put(f"{API}/complaints/{cid}", json={"assignees": []})
    r = session.get(f"{API}/complaints", params={"assigned_to": "unassigned"})
    assert r.status_code == 200
    ids = [x["id"] for x in r.json()["data"]]
    assert cid in ids

    # Cleanup
    r = session.delete(f"{API}/complaints/{cid}")
    assert r.status_code == 200


# ────────────────────── EXPORT / IMPORT ──────────────────────

def test_export_headers(session):
    r = session.get(f"{API}/complaints/export")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "Departments" in headers
    assert "Assignees" in headers
    assert "Category" not in headers


def test_import_template(session):
    r = session.get(f"{API}/complaints/import-template")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "Departments" in headers
    assert "Assignees" in headers
    assert "Category" not in headers


def test_import_upload(session):
    wb = Workbook()
    ws = wb.active
    ws.title = "Issues"
    ws.append([
        "Issue Number", "Title", "Details", "Customer", "Linked To", "City", "SKUs",
        "Departments", "Priority", "Status", "Assignees", "Created At",
    ])
    ws.append([
        "", "TEST_import_multi_v2", "imported", "Cust X", "", "",
        "", "Production, Quality", "high", "open", "Alice, Bob", "",
    ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # multipart — override content-type
    headers = {k: v for k, v in session.headers.items() if k.lower() != "content-type"}
    r = requests.post(
        f"{API}/complaints/import",
        headers=headers,
        files={"file": ("test.xlsx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        timeout=30,
    )
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 1

    # find the created issue
    r = session.get(f"{API}/complaints", params={"search": "TEST_import_multi_v2"})
    assert r.status_code == 200
    docs = r.json()["data"]
    assert docs
    doc = docs[0]
    assert set(doc["departments"]) == {"Production", "Quality"}
    assert len(doc["assignees"]) == 2
    assert {a["name"] for a in doc["assignees"]} == {"Alice", "Bob"}

    # cleanup
    session.delete(f"{API}/complaints/{doc['id']}")


def test_legacy_migrated_issue(session):
    r = session.get(f"{API}/complaints", params={"search": "CMP-2026-0001"})
    assert r.status_code == 200
    docs = r.json()["data"]
    assert docs, "expected pre-existing CMP-2026-0001 issue"
    doc = docs[0]
    # migrated: departments array should include Production
    assert "Production" in (doc.get("departments") or [])
