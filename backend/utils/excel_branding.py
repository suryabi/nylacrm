"""Apply Nyla brand styling (Helvetica font + logo header) to generated Excel workbooks."""
import os
from openpyxl.styles import Font
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils.cell import coordinate_from_string

BRAND_FONT = "Helvetica"
COMPANY_NAME = "Nyla Air Water"
_LOGO_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "nyla-logo.png")
_HEADER_ROWS = 4


def apply_brand_font(wb, name: str = BRAND_FONT):
    """Set the brand font on every cell, preserving all other font attributes."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                f = cell.font
                cell.font = Font(
                    name=name, size=f.size, bold=f.bold, italic=f.italic,
                    underline=f.underline, strike=f.strike, color=f.color,
                    vertAlign=f.vertAlign,
                )
    return wb


def add_brand_header(wb):
    """Insert a branded header (kingfisher logo + company name) at the top of each sheet."""
    for ws in wb.worksheets:
        ws.insert_rows(1, _HEADER_ROWS)
        if ws.freeze_panes:
            col, row = coordinate_from_string(ws.freeze_panes)
            ws.freeze_panes = f"{col}{row + _HEADER_ROWS}"
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 24
        ws.row_dimensions[3].height = 24
        ws.row_dimensions[4].height = 8
        if os.path.exists(_LOGO_PATH):
            img = XLImage(_LOGO_PATH)
            img.width = 68
            img.height = 68
            ws.add_image(img, "A1")
        name_cell = ws.cell(row=2, column=2, value=COMPANY_NAME)
        name_cell.font = Font(name=BRAND_FONT, size=14, bold=True, color="1D1D1F")
        tag_cell = ws.cell(row=3, column=2, value="Born from Air. Perfected by Nature.")
        tag_cell.font = Font(name=BRAND_FONT, size=9, color="86868B")
    return wb


def brand_workbook(wb):
    """Full branding pass: logo header + Helvetica on every cell."""
    add_brand_header(wb)
    apply_brand_font(wb)
    return wb
