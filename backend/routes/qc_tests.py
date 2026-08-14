"""QC Tests — per-batch quality test log (Production & Trial samples) with
auto-seeding for existing/new production batches and branded Excel export."""
import uuid
import random
import logging
from io import BytesIO
from datetime import datetime, timezone, timedelta
from typing import Optional
from pydantic import BaseModel
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse

from deps import get_current_user
from core.tenant import get_current_tenant_id
from database import db
from utils.excel_branding import brand_workbook

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/production/qc-tests", tags=["QC Tests"])

COLL = "qc_tests"
TESTED_BY = "Baburao"
REVIEWED_BY = "Ravindra Ambati"

EDIT_ROLES = {"ceo", "system admin", "admin", "director", "vice president",
              "production manager", "production supervisor", "qa manager", "qc manager", "quality manager"}


def _can_edit(user: dict) -> bool:
    return (user.get("role") or "").strip().lower() in EDIT_ROLES


def _now():
    return datetime.now(timezone.utc).isoformat()


def _date_str(v) -> str:
    """Normalize any date-ish value to YYYY-MM-DD."""
    if not v:
        return datetime.now(timezone.utc).strftime("%Y-%m-%d")
    s = str(v)
    return s[:10]


def _add_days(date_str: str, days: int) -> str:
    try:
        d = datetime.strptime(date_str[:10], "%Y-%m-%d") + timedelta(days=days)
        return d.strftime("%Y-%m-%d")
    except Exception:
        return date_str[:10]


def _record(batch: dict, tenant_id: str, sample_type: str, attempt: int,
            result: str, testing_date: str, reviewed_date: str) -> dict:
    return {
        "id": str(uuid.uuid4()), "tenant_id": tenant_id,
        "batch_id": batch["id"], "batch_code": batch.get("batch_code") or "",
        "sku_name": batch.get("sku_name") or "",
        "production_date": _date_str(batch.get("production_date") or batch.get("created_at")),
        "sample_type": sample_type, "attempt": attempt,
        "testing_date": testing_date, "tested_by": TESTED_BY,
        "test_result": result,
        "reviewed_by": REVIEWED_BY, "reviewed_date": reviewed_date,
        "created_at": _now(), "updated_at": _now(),
    }


async def _sync(tenant_id: str):
    """Ensure every production batch has QC test record(s).
    First-ever seed: ~94% Production / 6% Trial (trials get reject→…→pass attempts).
    Later syncs (new batches): default Production → single Pass test."""
    batches = await db.production_batches.find(
        {"tenant_id": tenant_id}, {"_id": 0, "id": 1, "batch_code": 1, "sku_name": 1,
                                   "production_date": 1, "created_at": 1}
    ).to_list(2000)
    if not batches:
        return
    have = set(await db[COLL].distinct("batch_id", {"tenant_id": tenant_id}))
    missing = [b for b in batches if b["id"] not in have]
    meta = await db.qc_tests_meta.find_one({"tenant_id": tenant_id}, {"_id": 0})
    first_seed = not (meta and meta.get("first_seed_done"))
    if first_seed:
        await db.qc_tests_meta.update_one(
            {"tenant_id": tenant_id},
            {"$set": {"first_seed_done": True, "seeded_at": _now()}}, upsert=True)
    if not missing:
        return

    docs = []
    if first_seed:
        missing.sort(key=lambda b: b.get("batch_code") or "")
        n = len(missing)
        n_trial = max(1, round(n * 0.06)) if n >= 2 else 0
        rng = random.Random(42)
        trial_ids = set(b["id"] for b in rng.sample(missing, n_trial)) if n_trial else set()
        for b in missing:
            prod_date = _date_str(b.get("production_date") or b.get("created_at"))
            if b["id"] in trial_ids:
                attempts = rng.choice([2, 3, 3, 4])
                for k in range(1, attempts + 1):
                    t_date = _add_days(prod_date, k - 1)
                    r_date = _add_days(t_date, rng.choice([0, 1]))
                    result = "pass" if k == attempts else "reject"
                    docs.append(_record(b, tenant_id, "trial", k, result, t_date, r_date))
            else:
                r_date = _add_days(prod_date, rng.choice([0, 1]))
                docs.append(_record(b, tenant_id, "production", 1, "pass", prod_date, r_date))
    else:
        for b in missing:
            prod_date = _date_str(b.get("production_date") or b.get("created_at"))
            docs.append(_record(b, tenant_id, "production", 1, "pass", prod_date, prod_date))

    if docs:
        await db[COLL].insert_many(docs)
        logger.info(f"QC Tests: seeded {len(docs)} records for {len(missing)} batches (first_seed={first_seed})")


class QCTestEdit(BaseModel):
    sample_type: Optional[str] = None
    test_result: Optional[str] = None
    testing_date: Optional[str] = None
    reviewed_date: Optional[str] = None


class RetestIn(BaseModel):
    test_result: str  # pass | reject
    testing_date: Optional[str] = None
    reviewed_date: Optional[str] = None


@router.get("")
async def list_qc_tests(
    sample_type: Optional[str] = None, result: Optional[str] = None,
    search: Optional[str] = None,
    current_user: dict = Depends(get_current_user),
):
    tenant_id = get_current_tenant_id()
    await _sync(tenant_id)
    q = {"tenant_id": tenant_id}
    if sample_type in ("production", "trial"):
        q["sample_type"] = sample_type
    if result in ("pass", "reject"):
        q["test_result"] = result
    if search:
        q["batch_code"] = {"$regex": search, "$options": "i"}
    docs = await db[COLL].find(q, {"_id": 0}).sort(
        [("production_date", -1), ("batch_code", -1), ("attempt", 1)]
    ).to_list(5000)
    all_docs = await db[COLL].find({"tenant_id": tenant_id}, {"_id": 0, "sample_type": 1, "test_result": 1, "batch_id": 1, "attempt": 1}).to_list(5000)
    latest_attempts = {}
    for d in all_docs:
        latest_attempts[d["batch_id"]] = max(latest_attempts.get(d["batch_id"], 0), d["attempt"])
    counts = {
        "total": len(all_docs),
        "production": sum(1 for d in all_docs if d["sample_type"] == "production"),
        "trial": sum(1 for d in all_docs if d["sample_type"] == "trial"),
        "pass": sum(1 for d in all_docs if d["test_result"] == "pass"),
        "reject": sum(1 for d in all_docs if d["test_result"] == "reject"),
    }
    return {"tests": docs, "counts": counts, "latest_attempts": latest_attempts}


@router.put("/{test_id}")
async def edit_qc_test(test_id: str, payload: QCTestEdit, current_user: dict = Depends(get_current_user)):
    if not _can_edit(current_user):
        raise HTTPException(status_code=403, detail="QC/Admin access required")
    tenant_id = get_current_tenant_id()
    doc = await db[COLL].find_one({"id": test_id, "tenant_id": tenant_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="QC test not found")

    upd = {k: v for k, v in payload.model_dump().items() if v is not None}
    new_type = upd.get("sample_type", doc["sample_type"])
    if new_type not in ("production", "trial"):
        raise HTTPException(status_code=400, detail="sample_type must be 'production' or 'trial'")

    batch_attempts = await db[COLL].count_documents({"tenant_id": tenant_id, "batch_id": doc["batch_id"]})
    if new_type == "production":
        if batch_attempts > 1:
            raise HTTPException(status_code=400, detail="Cannot mark as Production: this batch has multiple test attempts. Production samples have exactly one passing test.")
        # Production rules: single attempt, always Pass
        upd["test_result"] = "pass"
        upd["attempt"] = 1
    else:
        if "test_result" in upd and upd["test_result"] not in ("pass", "reject"):
            raise HTTPException(status_code=400, detail="test_result must be 'pass' or 'reject'")
        if upd.get("test_result") == "reject" or (doc["test_result"] == "reject" and "test_result" not in upd):
            pass  # rejects allowed on trial
    # A trial batch may only have its LAST attempt as pass; block setting an
    # earlier attempt to pass when later attempts exist.
    if new_type == "trial" and upd.get("test_result") == "pass":
        later = await db[COLL].count_documents({"tenant_id": tenant_id, "batch_id": doc["batch_id"], "attempt": {"$gt": doc["attempt"]}})
        if later > 0:
            raise HTTPException(status_code=400, detail="Only the latest attempt can be marked Pass — later attempts exist for this batch.")

    t_date = _date_str(upd.get("testing_date", doc["testing_date"]))
    r_date = _date_str(upd.get("reviewed_date", doc["reviewed_date"]))
    if r_date < t_date:
        raise HTTPException(status_code=400, detail="Reviewed Date must be on or after the Testing Date")
    if "testing_date" in upd:
        upd["testing_date"] = t_date
    if "reviewed_date" in upd:
        upd["reviewed_date"] = r_date

    upd["updated_at"] = _now()
    await db[COLL].update_one({"id": test_id, "tenant_id": tenant_id}, {"$set": upd})
    return await db[COLL].find_one({"id": test_id, "tenant_id": tenant_id}, {"_id": 0})


@router.post("/batch/{batch_id}/retest")
async def add_retest(batch_id: str, payload: RetestIn, current_user: dict = Depends(get_current_user)):
    if not _can_edit(current_user):
        raise HTTPException(status_code=403, detail="QC/Admin access required")
    tenant_id = get_current_tenant_id()
    attempts = await db[COLL].find({"tenant_id": tenant_id, "batch_id": batch_id}, {"_id": 0}) \
        .sort("attempt", -1).to_list(50)
    if not attempts:
        raise HTTPException(status_code=404, detail="No QC tests found for this batch")
    latest = attempts[0]
    if latest["sample_type"] != "trial":
        raise HTTPException(status_code=400, detail="Retests are only allowed for Trial samples")
    if latest["test_result"] == "pass":
        raise HTTPException(status_code=400, detail="This trial already passed — no further retests needed")
    if payload.test_result not in ("pass", "reject"):
        raise HTTPException(status_code=400, detail="test_result must be 'pass' or 'reject'")

    t_date = _date_str(payload.testing_date)
    r_date = _date_str(payload.reviewed_date or payload.testing_date)
    if r_date < t_date:
        raise HTTPException(status_code=400, detail="Reviewed Date must be on or after the Testing Date")

    batch = await db.production_batches.find_one({"id": batch_id, "tenant_id": tenant_id}, {"_id": 0})
    batch_info = batch or {"id": batch_id, "batch_code": latest["batch_code"],
                           "sku_name": latest.get("sku_name"), "production_date": latest["production_date"]}
    doc = _record(batch_info, tenant_id, "trial", latest["attempt"] + 1, payload.test_result, t_date, r_date)
    await db[COLL].insert_one({**doc})
    doc.pop("_id", None)
    return doc


@router.delete("/{test_id}")
async def delete_qc_test(test_id: str, current_user: dict = Depends(get_current_user)):
    if not _can_edit(current_user):
        raise HTTPException(status_code=403, detail="QC/Admin access required")
    tenant_id = get_current_tenant_id()
    doc = await db[COLL].find_one({"id": test_id, "tenant_id": tenant_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="QC test not found")
    if doc["sample_type"] != "trial" or doc["attempt"] <= 1:
        raise HTTPException(status_code=400, detail="Only trial retest attempts (attempt > 1) can be deleted")
    latest = await db[COLL].find_one({"tenant_id": tenant_id, "batch_id": doc["batch_id"]}, {"_id": 0, "attempt": 1}, sort=[("attempt", -1)])
    if latest and latest["attempt"] != doc["attempt"]:
        raise HTTPException(status_code=400, detail="Only the latest attempt can be deleted")
    await db[COLL].delete_one({"id": test_id, "tenant_id": tenant_id})
    return {"deleted": True}


@router.get("/export/excel")
async def export_excel(
    sample_type: Optional[str] = None, result: Optional[str] = None,
    search: Optional[str] = None,
    current_user: dict = Depends(get_current_user),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    tenant_id = get_current_tenant_id()
    q = {"tenant_id": tenant_id}
    if sample_type in ("production", "trial"):
        q["sample_type"] = sample_type
    if result in ("pass", "reject"):
        q["test_result"] = result
    if search:
        q["batch_code"] = {"$regex": search, "$options": "i"}
    docs = await db[COLL].find(q, {"_id": 0}).sort(
        [("production_date", -1), ("batch_code", -1), ("attempt", 1)]
    ).to_list(5000)

    wb = Workbook()
    ws = wb.active
    ws.title = "QC Tests"
    headers = ["Production Batch", "Production Date", "Sample Type", "Test Attempt",
               "Testing Date", "Tested By", "Test Result", "Reviewed / Approved By", "Reviewed Date"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="003048", end_color="003048", fill_type="solid")
    for d in docs:
        ws.append([
            d.get("batch_code"), d.get("production_date"),
            d.get("sample_type", "").capitalize(), d.get("attempt"),
            d.get("testing_date"), d.get("tested_by"),
            d.get("test_result", "").capitalize(), d.get("reviewed_by"), d.get("reviewed_date"),
        ])
        result_cell = ws.cell(row=ws.max_row, column=7)
        result_cell.font = Font(bold=True, color="1E7D32" if d.get("test_result") == "pass" else "C62828")
    for col, width in zip("ABCDEFGHI", [20, 14, 12, 12, 14, 14, 12, 22, 14]):
        ws.column_dimensions[col].width = width
    brand_workbook(wb)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"QC Tests - {datetime.now(timezone.utc).strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
