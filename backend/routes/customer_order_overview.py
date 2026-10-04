"""Customer Order Overview — completed regular Stock-Out bottles by month.

Accounts that share an active Brand Group Child SKU are collapsed into one
logical row so their delivered bottle volumes are counted once and summed.
"""

import io
from collections import defaultdict
from datetime import datetime, timezone

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from database import get_tenant_db
from deps import get_current_user
from utils.excel_branding import brand_workbook
from routes.customer_invoice_overview import _default_distributor, _distributor_cities, _month_keys


router = APIRouter()
MONTHS = 6
COMPLETED_STATUSES = ["delivered", "complete", "completed"]


class _DisjointSet:
    def __init__(self, ids: list[str]):
        self.parent = {item: item for item in ids}

    def find(self, item: str) -> str:
        while self.parent[item] != item:
            self.parent[item] = self.parent[self.parent[item]]
            item = self.parent[item]
        return item

    def union(self, items: list[str]):
        if len(items) < 2:
            return
        root = self.find(items[0])
        for item in items[1:]:
            other = self.find(item)
            if other != root:
                self.parent[other] = root


def _next_month(month_key: str) -> str:
    year, month = map(int, month_key.split("-"))
    return f"{year + (month == 12):04d}-{1 if month == 12 else month + 1:02d}-01"


def _month_of(delivery: dict) -> str:
    value = delivery.get("delivered_at") or delivery.get("delivery_date") or ""
    return str(value)[:7]


def _units(item: dict) -> int:
    try:
        return max(0, int(item.get("quantity") or 0))
    except (TypeError, ValueError):
        return 0


async def _account_distributors(tdb, account_ids: list[str]) -> dict:
    names = {
        row["id"]: row.get("distributor_name")
        async for row in tdb.distributors.find({}, {"_id": 0, "id": 1, "distributor_name": 1})
    }
    resolved = {}
    async for assignment in tdb.account_distributor_assignments.find(
        {"account_id": {"$in": account_ids}, "status": {"$ne": "inactive"}},
        {"_id": 0, "account_id": 1, "distributor_id": 1, "distributor_name": 1, "is_primary": 1},
    ).sort("is_primary", -1):
        resolved.setdefault(assignment["account_id"], {
            "id": assignment.get("distributor_id"),
            "name": names.get(assignment.get("distributor_id")) or assignment.get("distributor_name"),
        })

    latest = {}
    async for delivery in tdb.distributor_deliveries.find(
        {
            "account_id": {"$in": account_ids},
            "status": {"$in": COMPLETED_STATUSES},
            "is_promo": {"$ne": True},
        },
        {"_id": 0, "account_id": 1, "distributor_id": 1, "distributor_name": 1, "delivery_date": 1, "delivered_at": 1},
    ):
        stamp = delivery.get("delivered_at") or delivery.get("delivery_date") or ""
        account_id = delivery.get("account_id")
        if account_id and stamp > (latest.get(account_id) or {}).get("stamp", ""):
            latest[account_id] = {
                "stamp": stamp,
                "id": delivery.get("distributor_id"),
                "name": names.get(delivery.get("distributor_id")) or delivery.get("distributor_name"),
            }
    for account_id, distributor in latest.items():
        if distributor.get("id"):
            resolved[account_id] = {"id": distributor["id"], "name": distributor.get("name")}
    return resolved


async def _logical_rows(tdb, accounts: list[dict]) -> tuple[list[dict], dict[str, str]]:
    """Return deduplicated display rows and account → logical-row mapping."""
    account_by_id = {account["id"]: account for account in accounts}
    account_ids = list(account_by_id)
    groups = await tdb.brand_groups.find(
        {"is_active": {"$ne": False}, "account_ids": {"$in": account_ids}},
        {"_id": 0, "id": 1, "name": 1, "code": 1, "account_ids": 1},
    ).to_list(2000)
    group_by_id = {group["id"]: group for group in groups}
    shared_children = await tdb.child_skus.find(
        {"is_active": {"$ne": False}, "brand_group_id": {"$in": list(group_by_id)}},
        {"_id": 0, "id": 1, "child_sku_name": 1, "child_sku_code": 1, "brand_group_id": 1},
    ).to_list(5000)

    dsu = _DisjointSet(account_ids)
    child_members = {}
    for child in shared_children:
        group = group_by_id.get(child.get("brand_group_id")) or {}
        members = sorted(set(group.get("account_ids") or []) & set(account_ids))
        if len(members) > 1:
            dsu.union(members)
            child_members[child["id"]] = members

    components = defaultdict(list)
    for account_id in account_ids:
        components[dsu.find(account_id)].append(account_id)

    logical = []
    account_to_row = {}
    for root, member_ids in components.items():
        member_set = set(member_ids)
        children = [child for child in shared_children if member_set.intersection(child_members.get(child["id"], []))]
        group_ids = sorted({child.get("brand_group_id") for child in children if child.get("brand_group_id")})
        is_group = len(member_ids) > 1 and bool(children)
        if is_group:
            row_id = "group:" + ":".join(group_ids)
            group_names = sorted({group_by_id[group_id].get("name") or "Brand Group" for group_id in group_ids})
            group_codes = sorted({group_by_id[group_id].get("code") for group_id in group_ids if group_by_id[group_id].get("code")})
            child_names = sorted({child.get("child_sku_name") or "Shared Child SKU" for child in children})
            member_accounts = sorted((account_by_id[item] for item in member_ids), key=lambda row: (row.get("account_name") or "").lower())
            cities = sorted({account.get("city") for account in member_accounts if account.get("city")})
            logical.append({
                "account_id": row_id,
                "account_code": " / ".join(group_codes) or None,
                "account_name": " / ".join(group_names) or "Shared Child SKU Group",
                "city": ", ".join(cities),
                "row_type": "brand_group",
                "member_account_ids": member_ids,
                "member_account_names": [account.get("account_name") or "—" for account in member_accounts],
                "member_count": len(member_ids),
                "shared_child_sku_names": child_names,
            })
        else:
            account = account_by_id[root]
            row_id = account["id"]
            logical.append({
                "account_id": row_id,
                "account_code": account.get("account_id"),
                "account_name": account.get("account_name") or "—",
                "city": account.get("city"),
                "row_type": "account",
                "member_account_ids": [account["id"]],
                "member_account_names": [account.get("account_name") or "—"],
                "member_count": 1,
                "shared_child_sku_names": [],
            })
        for account_id in member_ids:
            account_to_row[account_id] = row_id
    return logical, account_to_row


async def _build(current_user: dict | None = None) -> dict:
    tdb = get_tenant_db()
    today = datetime.now(timezone.utc).date()
    all_keys = _month_keys(today, MONTHS + 1)
    prev_key, keys = all_keys[0], all_keys[1:]
    labels = [datetime.strptime(key, "%Y-%m").strftime("%b %Y") for key in keys]

    accounts = await tdb.accounts.find(
        {"status": {"$in": ["active", None]}},
        {"_id": 0, "id": 1, "account_id": 1, "account_name": 1, "city": 1},
    ).to_list(5000)
    account_ids = [account["id"] for account in accounts]
    rows, account_to_row = await _logical_rows(tdb, accounts)
    first_day = f"{prev_key}-01"
    next_day = _next_month(keys[-1])
    query = {
        "account_id": {"$in": account_ids},
        "status": {"$in": COMPLETED_STATUSES},
        "is_promo": {"$ne": True},
        "$or": [
            {"delivered_at": {"$gte": first_day, "$lt": next_day}},
            {
                "$and": [
                    {"$or": [{"delivered_at": None}, {"delivered_at": ""}, {"delivered_at": {"$exists": False}}]},
                    {"delivery_date": {"$gte": first_day, "$lt": next_day}},
                ]
            },
        ],
    }
    deliveries = await tdb.distributor_deliveries.find(
        query,
        {"_id": 0, "id": 1, "account_id": 1, "delivery_date": 1, "delivered_at": 1, "total_quantity": 1},
    ).to_list(100000)
    delivery_ids = [delivery["id"] for delivery in deliveries]
    items_by_delivery = defaultdict(list)
    if delivery_ids:
        async for item in tdb.distributor_delivery_items.find(
            {"delivery_id": {"$in": delivery_ids}},
            {"_id": 0, "delivery_id": 1, "quantity": 1},
        ):
            items_by_delivery[item["delivery_id"]].append(item)

    totals = defaultdict(lambda: defaultdict(int))
    counts = defaultdict(int)
    for delivery in deliveries:
        row_id = account_to_row.get(delivery.get("account_id"))
        month = _month_of(delivery)
        if not row_id or month not in all_keys:
            continue
        items = items_by_delivery.get(delivery["id"], [])
        bottles = sum(_units(item) for item in items)
        if not items:
            bottles = _units(delivery)
            if not bottles:
                try:
                    bottles = max(0, int(delivery.get("total_quantity") or 0))
                except (TypeError, ValueError):
                    bottles = 0
        totals[row_id][month] += bottles
        if month in keys:
            counts[row_id] += 1

    distributor_by_account = await _account_distributors(tdb, account_ids)
    distributor_names = {
        row["id"]: row.get("distributor_name")
        async for row in tdb.distributors.find({}, {"_id": 0, "id": 1, "distributor_name": 1})
    }
    for row in rows:
        values = [int(totals[row["account_id"]].get(key, 0)) for key in keys]
        previous = int(totals[row["account_id"]].get(prev_key, 0))
        member_distributors = [distributor_by_account.get(account_id) or {} for account_id in row["member_account_ids"]]
        distributor_ids = sorted({item.get("id") for item in member_distributors if item.get("id")})
        row.update({
            "distributor_ids": distributor_ids,
            "distributor_id": distributor_ids[0] if len(distributor_ids) == 1 else None,
            "distributor_name": ", ".join(distributor_names.get(item) or "Unknown" for item in distributor_ids) or "Unassigned",
            "has_unassigned": any(not item.get("id") for item in member_distributors),
            "months": values,
            "prev_month": previous,
            "total": sum(values),
            "mom_change": values[-1] - values[-2],
            "delivery_count": counts.get(row["account_id"], 0),
        })
    rows.sort(key=lambda row: (row["account_name"] or "").lower())

    month_totals = [sum(row["months"][index] for row in rows) for index in range(MONTHS)]
    previous_total = sum(row["prev_month"] for row in rows)
    used_distributors = sorted(
        {(item, distributor_names.get(item) or "Unknown") for row in rows for item in row["distributor_ids"]},
        key=lambda item: item[1].lower(),
    )
    dist_cities = await _distributor_cities(tdb)
    default_rows = [{**row, "distributor_id": item} for row in rows for item in row["distributor_ids"]]
    return {
        "months": [{"key": key, "label": label, "current": key == keys[-1]} for key, label in zip(keys, labels)],
        "rows": rows,
        "month_totals": month_totals,
        "prev_month_total": previous_total,
        "mom_change_total": month_totals[-1] - month_totals[-2],
        "grand_total": sum(month_totals),
        "default_distributor_id": _default_distributor(current_user, default_rows, dist_cities) if current_user else None,
        "filters": {
            "distributors": [{"id": item, "name": name} for item, name in used_distributors]
            + ([{"id": "unassigned", "name": "Unassigned"}] if any(row["has_unassigned"] for row in rows) else [])
        },
        "summary": {
            "rows": len(rows),
            "source_accounts": len(accounts),
            "grouped_rows": sum(1 for row in rows if row["row_type"] == "brand_group"),
            "ordered_rows": sum(1 for row in rows if row["total"] > 0),
            "deliveries": sum(counts.values()),
        },
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


@router.get("/reports/customer-order-overview")
async def customer_order_overview(current_user: dict = Depends(get_current_user)):
    return await _build(current_user)


@router.get("/reports/customer-order-overview/export")
async def export_customer_order_overview(current_user: dict = Depends(get_current_user)):
    data = await _build()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Customer Order Overview"
    sheet.append(["Customer Order Overview — completed Stock-Out bottles, last 6 months"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([f"Generated: {datetime.now(timezone.utc).strftime('%d %b %Y %H:%M UTC')}"])
    sheet.append([])
    sheet.append([
        "Customer / Brand Group", "Row Type", "Shared Child SKU(s)", "Member Accounts", "Code", "City", "Distributor",
        *[month["label"] for month in data["months"]], "Change vs last month",
    ])
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in data["rows"]:
        sheet.append([
            row["account_name"], "Brand Group / Shared SKU" if row["row_type"] == "brand_group" else "Account",
            ", ".join(row["shared_child_sku_names"]), ", ".join(row["member_account_names"]),
            row.get("account_code") or "", row.get("city") or "", row.get("distributor_name") or "",
            *row["months"], row["mom_change"],
        ])
    sheet.append(["Total", "", "", "", "", "", "", *data["month_totals"], data["mom_change_total"]])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    for row in sheet.iter_rows(min_row=header_row + 1, min_col=8, max_col=8 + MONTHS):
        for cell in row:
            cell.number_format = "#,##0"
    widths = [34, 24, 42, 48, 18, 24, 28] + [14] * MONTHS + [18]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=header_row, column=index).column_letter].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=2)
    brand_workbook(workbook)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"customer-order-overview-{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )