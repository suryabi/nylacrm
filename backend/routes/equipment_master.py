"""Equipment Master & Maintenance Management — production equipment hierarchy,
maintenance/replacement schedules and vendor associations."""
import re
import uuid
import logging
from datetime import datetime, timedelta, timezone
from typing import Optional, List
from pydantic import BaseModel
from fastapi import APIRouter, HTTPException, Depends, Query
from deps import get_current_user, get_user_or_factory_staff
from core.tenant import get_current_tenant_id
from database import db

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/production/equipment", tags=["Equipment Master"])

DEFAULT_TYPES = ["Filters", "Pumps", "Air Water Generators", "RO/Membrane Systems",
                 "UV Systems", "Ozone Systems", "Filling Machines", "Compressors", "Chillers"]


class EquipmentTypeIn(BaseModel):
    name: str
    description: Optional[str] = None


class TypeCodeIn(BaseModel):
    code: str


class VendorAssociation(BaseModel):
    vendor_id: str
    vendor_name: Optional[str] = None
    preferred: bool = False
    lead_time_days: Optional[int] = None
    last_purchase_price: Optional[float] = None
    vendor_part_code: Optional[str] = None
    contact_details: Optional[str] = None
    amc_details: Optional[str] = None


class EquipmentIn(BaseModel):
    name: str
    equipment_type_id: str
    make_brand: Optional[str] = None
    model_spec: Optional[str] = None
    maintenance_frequency_days: Optional[int] = None
    replacement_frequency_days: Optional[int] = None
    maintenance_start_date: Optional[str] = None   # YYYY-MM-DD
    replacement_start_date: Optional[str] = None
    photo_inspections_per_month: Optional[int] = None
    procurement_lead_time_days: Optional[int] = None
    last_maintenance_date: Optional[str] = None   # YYYY-MM-DD
    next_maintenance_due: Optional[str] = None
    last_replacement_date: Optional[str] = None
    next_replacement_due: Optional[str] = None
    expected_life: Optional[str] = None
    min_stock_level: Optional[int] = None
    remarks: Optional[str] = None
    vendors: List[VendorAssociation] = []
    parent_equipment_id: Optional[str] = None
    has_children: bool = False
    facility_id: Optional[str] = None
    floor_id: Optional[str] = None
    area_id: Optional[str] = None
    per_run_maintenance: bool = False          # must be cleared before every production run
    per_run_sku_ids: List[str] = []            # empty = applies to every SKU
    post_run_maintenance: bool = False         # cleanup/maintenance required after every production run
    post_run_sku_ids: List[str] = []


class LogDateBody(BaseModel):
    date: Optional[str] = None  # YYYY-MM-DD, default today
    remarks: Optional[str] = None
    photo_ids: List[str] = []


async def _log_photos_meta(tenant_id: str, equipment_id: str, photo_ids: List[str]) -> list:
    """Resolve uploaded attachment ids into embedded log-photo references."""
    if not photo_ids:
        return []
    eq = await db.equipment_master.find_one({"tenant_id": tenant_id, "id": equipment_id}, {"_id": 0, "attachments": 1})
    by_id = {f["id"]: f for f in (eq or {}).get("attachments", [])}
    return [{"id": pid, "filename": by_id[pid]["filename"]} for pid in photo_ids if pid in by_id]


def _add_days(date_str: Optional[str], days: Optional[int]) -> Optional[str]:
    if not date_str or not days:
        return None
    try:
        return (datetime.fromisoformat(date_str[:10]) + timedelta(days=days)).strftime("%Y-%m-%d")
    except ValueError:
        return None


def _compute_dues(doc: dict):
    """Fill next-due dates from (last date or schedule start date) + frequency when not explicitly set."""
    if not doc.get("next_maintenance_due"):
        base = doc.get("last_maintenance_date") or doc.get("maintenance_start_date")
        doc["next_maintenance_due"] = _add_days(base, doc.get("maintenance_frequency_days"))
    if not doc.get("next_replacement_due"):
        base = doc.get("last_replacement_date") or doc.get("replacement_start_date")
        doc["next_replacement_due"] = _add_days(base, doc.get("replacement_frequency_days"))


# ── Equipment Types (parent master) ──────────────────────────

@router.get("/types")
async def list_types(current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    await _ensure_codes(tenant_id)
    types = await db.equipment_types.find({"tenant_id": tenant_id}, {"_id": 0}).sort("name", 1).to_list(500)
    counts = {}
    for t in types:
        counts[t["id"]] = await db.equipment_master.count_documents({"tenant_id": tenant_id, "equipment_type_id": t["id"]})
        t["equipment_count"] = counts[t["id"]]
    return {"types": types}


@router.post("/types")
async def create_type(data: EquipmentTypeIn, current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    if not data.name.strip():
        raise HTTPException(400, "Type name is required")
    dup = await db.equipment_types.find_one({"tenant_id": tenant_id, "name": {"$regex": f"^{data.name.strip()}$", "$options": "i"}})
    if dup:
        raise HTTPException(409, f"Equipment type '{data.name}' already exists")
    doc = {"id": str(uuid.uuid4()), "tenant_id": tenant_id, "name": data.name.strip(),
           "description": data.description, "created_at": datetime.now(timezone.utc).isoformat()}
    await db.equipment_types.insert_one(doc)
    doc.pop("_id", None)
    return doc


@router.post("/types/seed-defaults")
async def seed_default_types(current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    existing = {t["name"].lower() for t in await db.equipment_types.find({"tenant_id": tenant_id}, {"_id": 0, "name": 1}).to_list(500)}
    created = 0
    for name in DEFAULT_TYPES:
        if name.lower() not in existing:
            await db.equipment_types.insert_one({"id": str(uuid.uuid4()), "tenant_id": tenant_id, "name": name,
                                                 "description": None, "created_at": datetime.now(timezone.utc).isoformat()})
            created += 1
    return {"created": created}


@router.put("/types/{type_id}")
async def update_type(type_id: str, data: EquipmentTypeIn, current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    res = await db.equipment_types.update_one(
        {"tenant_id": tenant_id, "id": type_id},
        {"$set": {"name": data.name.strip(), "description": data.description}})
    if res.matched_count == 0:
        raise HTTPException(404, "Type not found")
    await db.equipment_master.update_many(
        {"tenant_id": tenant_id, "equipment_type_id": type_id},
        {"$set": {"equipment_type_name": data.name.strip()}})
    return await db.equipment_types.find_one({"tenant_id": tenant_id, "id": type_id}, {"_id": 0})


@router.put("/types/{type_id}/code")
async def set_type_code(type_id: str, data: TypeCodeIn, current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    code = re.sub(r"[^A-Za-z0-9]", "", data.code or "").upper()[:6]
    if not code:
        raise HTTPException(400, "Code must contain 1-6 letters/numbers")
    dup = await db.equipment_types.find_one({"tenant_id": tenant_id, "code": code, "id": {"$ne": type_id}}, {"_id": 0, "id": 1})
    if dup:
        raise HTTPException(400, f"Code {code} is already used by another equipment type")
    res = await db.equipment_types.update_one({"tenant_id": tenant_id, "id": type_id}, {"$set": {"code": code}})
    if res.matched_count == 0:
        raise HTTPException(404, "Type not found")
    return {"ok": True, "code": code}


@router.delete("/types/{type_id}")
async def delete_type(type_id: str, current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    n = await db.equipment_master.count_documents({"tenant_id": tenant_id, "equipment_type_id": type_id})
    if n > 0:
        raise HTTPException(409, f"Cannot delete: {n} equipment item(s) belong to this type")
    res = await db.equipment_types.delete_one({"tenant_id": tenant_id, "id": type_id})
    if res.deleted_count == 0:
        raise HTTPException(404, "Type not found")
    return {"deleted": True}


# ── Equipment / Components (child master) ────────────────────

async def _resolve_location(tenant_id: str, facility_id, floor_id, area_id) -> dict:
    """Validate facility → floor → area chain and return display names."""
    out = {"facility_name": None, "floor_name": None, "area_name": None}
    if not facility_id:
        return out
    fac = await db.facilities.find_one({"tenant_id": tenant_id, "id": facility_id},
                                       {"_id": 0, "facility_name": 1, "floors": 1})
    if not fac:
        raise HTTPException(400, "Facility not found")
    out["facility_name"] = fac["facility_name"]
    if floor_id:
        floor = next((f for f in fac.get("floors") or [] if f.get("id") == floor_id), None)
        if not floor:
            raise HTTPException(400, "Floor not found in this facility")
        out["floor_name"] = floor.get("name")
        if area_id:
            area = next((a for a in floor.get("areas") or [] if a.get("id") == area_id), None)
            if not area:
                raise HTTPException(400, "Area not found on this floor")
            out["area_name"] = area.get("name")
    return out


async def _validate_parent(tenant_id: str, parent_id: str, self_id: Optional[str] = None) -> dict:
    """Ensure parent exists and the chain has no cycle. Returns parent doc."""
    if self_id and parent_id == self_id:
        raise HTTPException(400, "Equipment cannot be its own parent")
    parent = await db.equipment_master.find_one({"tenant_id": tenant_id, "id": parent_id},
                                                {"_id": 0, "id": 1, "name": 1, "parent_equipment_id": 1})
    if not parent:
        raise HTTPException(400, "Parent equipment not found")
    seen, cur = set(), parent
    while cur and cur.get("parent_equipment_id"):
        pid = cur["parent_equipment_id"]
        if pid == self_id or pid in seen:
            raise HTTPException(400, "Invalid parent — this would create a circular hierarchy")
        seen.add(pid)
        cur = await db.equipment_master.find_one({"tenant_id": tenant_id, "id": pid},
                                                 {"_id": 0, "id": 1, "parent_equipment_id": 1})
    return parent


async def _hydrate_vendors(tenant_id: str, vendors: List[VendorAssociation]) -> list:
    out = []
    for v in vendors:
        d = v.dict()
        if not d.get("vendor_name"):
            vend = await db.inventory_vendors.find_one({"tenant_id": tenant_id, "id": v.vendor_id}, {"_id": 0, "vendor_name": 1})
            d["vendor_name"] = (vend or {}).get("vendor_name") or v.vendor_id
        out.append(d)
    return out


@router.get("")
async def list_equipment(
    equipment_type_id: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    current_user: dict = Depends(get_current_user),
):
    _ = current_user
    tenant_id = get_current_tenant_id()
    q = {"tenant_id": tenant_id}
    if equipment_type_id:
        q["equipment_type_id"] = equipment_type_id
    if search:
        q["$or"] = [{"name": {"$regex": search, "$options": "i"}},
                    {"make_brand": {"$regex": search, "$options": "i"}},
                    {"model_spec": {"$regex": search, "$options": "i"}}]
    items = await db.equipment_master.find(q, {"_id": 0}).sort([("equipment_type_name", 1), ("name", 1)]).to_list(2000)
    counts = {}
    async for row in db.equipment_master.aggregate([
        {"$match": {"tenant_id": tenant_id, "parent_equipment_id": {"$nin": [None, ""]}}},
        {"$group": {"_id": "$parent_equipment_id", "n": {"$sum": 1}}},
    ]):
        counts[row["_id"]] = row["n"]
    for it in items:
        it["children_count"] = counts.get(it["id"], 0)
    return {"equipment": items}


@router.post("")
async def create_equipment(data: EquipmentIn, current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    if not data.name.strip():
        raise HTTPException(400, "Equipment name is required")
    etype = await db.equipment_types.find_one({"tenant_id": tenant_id, "id": data.equipment_type_id}, {"_id": 0, "name": 1})
    if not etype:
        raise HTTPException(400, "Equipment type not found")
    doc = data.dict()
    doc["parent_equipment_name"] = None
    if data.parent_equipment_id:
        parent = await _validate_parent(tenant_id, data.parent_equipment_id)
        doc["parent_equipment_name"] = parent["name"]
    doc.update(await _resolve_location(tenant_id, data.facility_id, data.floor_id, data.area_id))
    if not data.facility_id:
        doc["floor_id"] = doc["area_id"] = None
    doc["vendors"] = await _hydrate_vendors(tenant_id, data.vendors)
    _compute_dues(doc)
    doc.update({
        "id": str(uuid.uuid4()), "tenant_id": tenant_id,
        "equipment_type_name": etype["name"], "is_active": True,
        "maintenance_log": [], "replacement_log": [],
        "created_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    })
    await db.equipment_master.insert_one(doc)
    doc.pop("_id", None)
    return doc


@router.put("/{equipment_id}")
async def update_equipment(equipment_id: str, data: EquipmentIn, current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    existing = await db.equipment_master.find_one({"tenant_id": tenant_id, "id": equipment_id}, {"_id": 0})
    if not existing:
        raise HTTPException(404, "Equipment not found")
    etype = await db.equipment_types.find_one({"tenant_id": tenant_id, "id": data.equipment_type_id}, {"_id": 0, "name": 1})
    if not etype:
        raise HTTPException(400, "Equipment type not found")
    updates = data.dict()
    updates["parent_equipment_name"] = None
    if data.parent_equipment_id:
        parent = await _validate_parent(tenant_id, data.parent_equipment_id, self_id=equipment_id)
        updates["parent_equipment_name"] = parent["name"]
    updates.update(await _resolve_location(tenant_id, data.facility_id, data.floor_id, data.area_id))
    if not data.facility_id:
        updates["floor_id"] = updates["area_id"] = None
    updates["vendors"] = await _hydrate_vendors(tenant_id, data.vendors)
    _compute_dues(updates)
    updates["equipment_type_name"] = etype["name"]
    updates["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.equipment_master.update_one({"tenant_id": tenant_id, "id": equipment_id}, {"$set": updates})
    if data.name.strip() != existing.get("name"):
        await db.equipment_master.update_many(
            {"tenant_id": tenant_id, "parent_equipment_id": equipment_id},
            {"$set": {"parent_equipment_name": data.name.strip()}})
    return await db.equipment_master.find_one({"tenant_id": tenant_id, "id": equipment_id}, {"_id": 0})


@router.delete("/{equipment_id}")
async def delete_equipment(equipment_id: str, current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    n = await db.equipment_master.count_documents({"tenant_id": tenant_id, "parent_equipment_id": equipment_id})
    if n > 0:
        raise HTTPException(409, f"Cannot delete: {n} child equipment item(s) belong to this equipment. Reassign or delete them first.")
    res = await db.equipment_master.delete_one({"tenant_id": tenant_id, "id": equipment_id})
    if res.deleted_count == 0:
        raise HTTPException(404, "Equipment not found")
    return {"deleted": True}


@router.post("/{equipment_id}/log-maintenance")
async def log_maintenance(equipment_id: str, body: LogDateBody, current_user: dict = Depends(get_current_user)):
    tenant_id = get_current_tenant_id()
    eq = await db.equipment_master.find_one({"tenant_id": tenant_id, "id": equipment_id}, {"_id": 0})
    if not eq:
        raise HTTPException(404, "Equipment not found")
    date = body.date or datetime.now(timezone.utc).strftime("%Y-%m-%d")
    next_due = _add_days(date, eq.get("maintenance_frequency_days"))
    entry = {"date": date, "remarks": body.remarks, "by": current_user.get("name") or current_user.get("email"),
             "photos": await _log_photos_meta(tenant_id, equipment_id, body.photo_ids),
             "at": datetime.now(timezone.utc).isoformat()}
    await db.equipment_master.update_one(
        {"tenant_id": tenant_id, "id": equipment_id},
        {"$set": {"last_maintenance_date": date, "next_maintenance_due": next_due,
                  "updated_at": datetime.now(timezone.utc).isoformat()},
         "$push": {"maintenance_log": entry}})
    return await db.equipment_master.find_one({"tenant_id": tenant_id, "id": equipment_id}, {"_id": 0})


@router.post("/{equipment_id}/log-replacement")
async def log_replacement(equipment_id: str, body: LogDateBody, current_user: dict = Depends(get_current_user)):
    tenant_id = get_current_tenant_id()
    eq = await db.equipment_master.find_one({"tenant_id": tenant_id, "id": equipment_id}, {"_id": 0})
    if not eq:
        raise HTTPException(404, "Equipment not found")
    date = body.date or datetime.now(timezone.utc).strftime("%Y-%m-%d")
    next_due = _add_days(date, eq.get("replacement_frequency_days"))
    entry = {"date": date, "remarks": body.remarks, "by": current_user.get("name") or current_user.get("email"),
             "photos": await _log_photos_meta(tenant_id, equipment_id, body.photo_ids),
             "at": datetime.now(timezone.utc).isoformat()}
    await db.equipment_master.update_one(
        {"tenant_id": tenant_id, "id": equipment_id},
        {"$set": {"last_replacement_date": date, "next_replacement_due": next_due,
                  "updated_at": datetime.now(timezone.utc).isoformat()},
         "$push": {"replacement_log": entry}})
    return await db.equipment_master.find_one({"tenant_id": tenant_id, "id": equipment_id}, {"_id": 0})


# ── Attachments (photos, certifications, quotations, documents) ──
import os
from fastapi import UploadFile, File, Form
from fastapi.responses import FileResponse, Response
from services.object_storage import build_path, put_object, get_object, delete_object

EQ_UPLOAD_DIR = "/app/backend/uploads/equipment_docs"
EQ_MAX_FILE_MB = 15
EQ_CATEGORIES = {"photo", "certification", "quotation", "document", "instruction"}
EQ_TYPE_PREFIX = {"photo": "PHOT", "certification": "CERT", "quotation": "VEND", "document": "OTHE"}


def _eq_file_basename(equipment_name: str, category: str) -> str:
    """TYPE_Equipment_Name (first 3 words if longer), spaces → underscores."""
    import re as _re
    words = (equipment_name or "equipment").strip().split()
    if len(words) > 3:
        words = words[:3]
    name_part = _re.sub(r"[^A-Za-z0-9_]", "", "_".join(words).replace(" ", "_"))
    return f"{EQ_TYPE_PREFIX[category]}_{name_part}"


def _next_serial(attachments: list, base: str) -> int:
    """Per equipment + doc-type serial: max existing serial for this base + 1."""
    import re as _re
    top = 0
    for f in attachments or []:
        m = _re.match(rf"^{_re.escape(base)}_(\d+)\.", f.get("filename") or "")
        if m:
            top = max(top, int(m.group(1)))
    return top + 1


async def _eq_read_file(meta: dict):
    """Object storage first; legacy pod-local file fallback."""
    if meta.get("storage_path"):
        return await get_object(meta["storage_path"])
    path = os.path.join(EQ_UPLOAD_DIR, f"{meta['id']}_{meta['filename']}")
    if not os.path.exists(path):
        raise HTTPException(404, "File missing from storage")
    with open(path, "rb") as fh:
        return fh.read(), meta.get("content_type") or "application/octet-stream"


@router.post("/{equipment_id}/files")
async def upload_equipment_file(equipment_id: str, file: UploadFile = File(...),
                                category: str = Form("document"),
                                current_user: dict = Depends(get_current_user)):
    tenant_id = get_current_tenant_id()
    eq = await db.equipment_master.find_one({"tenant_id": tenant_id, "id": equipment_id},
                                            {"_id": 0, "id": 1, "name": 1, "attachments": 1})
    if not eq:
        raise HTTPException(404, "Equipment not found")
    if category not in EQ_CATEGORIES:
        raise HTTPException(400, f"Invalid category — use one of {sorted(EQ_CATEGORIES)}")
    content = await file.read()
    if len(content) > EQ_MAX_FILE_MB * 1024 * 1024:
        raise HTTPException(400, f"File too large (max {EQ_MAX_FILE_MB} MB)")
    file_id = str(uuid.uuid4())
    original_name = os.path.basename(file.filename or "file")
    ext = os.path.splitext(original_name)[1].lower()
    base = _eq_file_basename(eq.get("name"), category)
    serial = _next_serial(eq.get("attachments"), base)
    safe_name = f"{base}_{serial:02d}{ext}"
    storage_path = build_path(tenant_id, "equipment_docs", safe_name)
    await put_object(storage_path, content, file.content_type or "application/octet-stream")
    meta = {"id": file_id, "filename": safe_name, "original_filename": original_name,
            "size": len(content), "storage_path": storage_path,
            "content_type": file.content_type, "category": category,
            "uploaded_by": current_user.get("name") or current_user.get("email"),
            "uploaded_at": datetime.now(timezone.utc).isoformat()}
    await db.equipment_master.update_one(
        {"tenant_id": tenant_id, "id": equipment_id},
        {"$push": {"attachments": meta},
         "$set": {"updated_at": datetime.now(timezone.utc).isoformat()}})
    return meta


@router.get("/{equipment_id}/files/{file_id}")
async def download_equipment_file(equipment_id: str, file_id: str,
                                  current_user: dict = Depends(get_user_or_factory_staff)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    eq = await db.equipment_master.find_one({"tenant_id": tenant_id, "id": equipment_id}, {"_id": 0, "attachments": 1})
    if not eq:
        raise HTTPException(404, "Equipment not found")
    meta = next((f for f in eq.get("attachments", []) if f["id"] == file_id), None)
    if not meta:
        raise HTTPException(404, "File not found")
    data, ctype = await _eq_read_file(meta)
    return Response(data, media_type=meta.get("content_type") or ctype,
                    headers={"Content-Disposition": f'attachment; filename="{meta["filename"]}"'})


@router.delete("/{equipment_id}/files/{file_id}")
async def delete_equipment_file(equipment_id: str, file_id: str,
                                current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    eq = await db.equipment_master.find_one({"tenant_id": tenant_id, "id": equipment_id}, {"_id": 0, "attachments": 1})
    if not eq:
        raise HTTPException(404, "Equipment not found")
    meta = next((f for f in eq.get("attachments", []) if f["id"] == file_id), None)
    if not meta:
        raise HTTPException(404, "File not found")
    if meta.get("storage_path"):
        try:
            await delete_object(meta["storage_path"])
        except Exception:
            pass
    else:
        path = os.path.join(EQ_UPLOAD_DIR, f"{file_id}_{meta['filename']}")
        if os.path.exists(path):
            os.remove(path)
    await db.equipment_master.update_one(
        {"tenant_id": tenant_id, "id": equipment_id},
        {"$pull": {"attachments": {"id": file_id}},
         "$set": {"updated_at": datetime.now(timezone.utc).isoformat()}})
    return {"deleted": True}


@router.get("/{equipment_id}/files-zip")
async def download_equipment_zip(equipment_id: str, current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    eq = await db.equipment_master.find_one({"tenant_id": tenant_id, "id": equipment_id},
                                            {"_id": 0, "name": 1, "attachments": 1})
    if not eq:
        raise HTTPException(404, "Equipment not found")
    attachments = eq.get("attachments", [])
    if not attachments:
        raise HTTPException(404, "No documents to export")
    import io
    import zipfile
    folder_map = {"photo": "Photos", "certification": "Certifications",
                  "quotation": "Vendor_Documents", "document": "Other_Documents"}
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in attachments:
            try:
                data, _ = await _eq_read_file(f)
            except Exception:
                continue
            zf.writestr(f"{folder_map.get(f.get('category'), 'Other_Documents')}/{f['filename']}", data)
    import re as _re
    zip_name = _re.sub(r"[^A-Za-z0-9_]", "", (eq.get("name") or "equipment").replace(" ", "_")) + "_documents.zip"
    return Response(buf.getvalue(), media_type="application/zip",
                    headers={"Content-Disposition": f'attachment; filename="{zip_name}"'})


# ── Excel template / export / bulk import ─────────────────────────────────────
import io as _io
import re as _rex
import openpyxl
from openpyxl.styles import Font, PatternFill

XLSX_COLUMNS = [
    "Name*", "Equipment Type* (name or code)", "Make/Brand", "Model/Specification",
    "Parent Equipment", "Has Child Equipment (Yes/No)",
    "Facility (name or code)", "Floor (name or code)", "Area (name or code)",
    "Maintenance Frequency", "Maintenance Start Date", "Last Maintenance Date",
    "Replacement Frequency", "Replacement Start Date", "Last Replacement Date",
    "Procurement Lead Time (days)", "Expected Life", "Min Stock Level",
    "Photo Inspections per Month", "Remarks",
]


def _mk_code(name: str, taken: set, maxlen: int = 6) -> str:
    """Intuitive short code: initials for multi-word names, first 4 letters otherwise."""
    words = _rex.findall(r"[A-Za-z0-9]+", (name or "").upper())
    if not words:
        base = "X"
    elif len(words) == 1:
        base = words[0][:4]
    else:
        base = "".join(w[0] for w in words)[:maxlen]
    base = base[:maxlen] or "X"
    code, i = base, 2
    while code in taken:
        code = f"{base}{i}"
        i += 1
    taken.add(code)
    return code


async def _ensure_codes(tenant_id: str):
    """Assign persistent short codes to equipment types, facilities, floors and areas."""
    types = await db.equipment_types.find({"tenant_id": tenant_id}, {"_id": 0}).to_list(500)
    taken = {t["code"] for t in types if t.get("code")}
    for t in types:
        if not t.get("code"):
            t["code"] = _mk_code(t["name"], taken)
            await db.equipment_types.update_one({"tenant_id": tenant_id, "id": t["id"]}, {"$set": {"code": t["code"]}})
    facs = await db.facilities.find({"tenant_id": tenant_id}, {"_id": 0}).to_list(500)
    fac_taken = {f["code"] for f in facs if f.get("code")}
    for f in facs:
        changed = False
        if not f.get("code"):
            f["code"] = _mk_code(f["facility_name"], fac_taken)
            changed = True
        fl_taken = {fl.get("code") for fl in f.get("floors") or [] if fl.get("code")}
        for fl in f.get("floors") or []:
            if not fl.get("code"):
                fl["code"] = _mk_code(fl["name"], fl_taken)
                changed = True
            ar_taken = {a.get("code") for a in fl.get("areas") or [] if a.get("code")}
            for a in fl.get("areas") or []:
                if not a.get("code"):
                    a["code"] = _mk_code(a["name"], ar_taken)
                    changed = True
        if changed:
            await db.facilities.update_one({"tenant_id": tenant_id, "id": f["id"]},
                                           {"$set": {"code": f["code"], "floors": f.get("floors") or []}})
    return types, facs
FREQ_TO_LABEL = {1: "Daily", 2: "Alternate Days", 7: "Weekly", 10: "Every 10 days",
                 14: "Bi-Weekly", 30: "Monthly", 90: "Quarterly"}
LABEL_TO_FREQ = {"daily": 1, "alternate days": 2, "weekly": 7, "every 10 days": 10,
                 "bi-weekly": 14, "biweekly": 14, "fortnightly": 14, "monthly": 30, "quarterly": 90,
                 # short codes
                 "d": 1, "ad": 2, "w": 7, "e10": 10, "bw": 14, "fn": 14, "m": 30, "q": 90}
FREQ_CODES = [("D", "Daily"), ("AD", "Alternate Days"), ("W", "Weekly"), ("E10", "Every 10 days"),
              ("BW", "Bi-Weekly / Fortnightly (FN also works)"), ("M", "Monthly"), ("Q", "Quarterly")]


def _freq_out(days):
    if not days:
        return ""
    return FREQ_TO_LABEL.get(days, str(days))


def _freq_in(val):
    if val is None or val == "":
        return None
    s = str(val).strip().lower()
    if s in LABEL_TO_FREQ:
        return LABEL_TO_FREQ[s]
    try:
        n = int(float(s))
        return n if n > 0 else None
    except ValueError:
        raise ValueError(f"Invalid frequency '{val}' — use Daily/Alternate Days/Weekly/Every 10 days/Bi-Weekly/Monthly/Quarterly or a number of days")


def _date_in(val):
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if hasattr(val, "isoformat") and not isinstance(val, str):
        return val.isoformat()[:10]
    s = str(val).strip()[:10]
    datetime.strptime(s, "%Y-%m-%d")
    return s


def _int_in(val, field):
    if val is None or val == "":
        return None
    try:
        return int(float(str(val).strip()))
    except ValueError:
        raise ValueError(f"Invalid number for {field}: '{val}'")


def _xlsx_response(wb, filename):
    buf = _io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


def _styled_header(ws):
    fill = PatternFill(start_color="2E5A8F", end_color="2E5A8F", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    for i, col in enumerate(XLSX_COLUMNS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(16, len(col) + 4)


@router.get("/import-template")
async def equipment_import_template(current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    types, facs = await _ensure_codes(tenant_id)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Equipment"
    ws.append(XLSX_COLUMNS)
    _styled_header(ws)
    ex_type = types[0]["code"] if types else "MACH"
    ex_fac = facs[0]["code"] if facs else ""
    ws.append(["RO Machine 1", ex_type, "Pentair", "RO-5000", "", "Yes",
               ex_fac, "GF", "RFC",
               "W", "2026-06-01", "", "Q", "2026-06-01", "",
               30, "5 years", 1, 4, "Example row — delete before uploading"])
    ref = wb.create_sheet("Reference")
    bold = Font(bold=True)
    ref.append(["EQUIPMENT TYPE CODES", ""]); ref["A1"].font = bold
    ref.append(["Code", "Equipment Type"])
    for t in types:
        ref.append([t.get("code"), t["name"]])
    ref.append([])
    ref.append(["FREQUENCY CODES", ""]); ref[f"A{ref.max_row}"].font = bold
    ref.append(["Code", "Frequency"])
    for code, label in FREQ_CODES:
        ref.append([code, label])
    ref.append(["", "…or type any number of days"])
    ref.append([])
    ref.append(["LOCATION CODES", ""]); ref[f"A{ref.max_row}"].font = bold
    ref.append(["Facility Code", "Facility", "Floor Code", "Floor", "Area Codes"])
    for f in facs:
        floors = f.get("floors") or []
        if not floors:
            ref.append([f.get("code"), f["facility_name"], "", "", ""])
        for fl in floors:
            areas = ", ".join(f"{a.get('code')} = {a['name']}" for a in fl.get("areas") or []) or "-"
            ref.append([f.get("code"), f["facility_name"], fl.get("code"), fl.get("name"), areas])
    ref.append([])
    ref.append(["Notes:", "Use the short CODES above (or the full names — both work). Dates must be YYYY-MM-DD. Parent Equipment is matched by exact name (row order doesn't matter). Re-uploading a row with an existing Name updates that equipment."])
    for col, w in (("A", 16), ("B", 30), ("C", 12), ("D", 22), ("E", 50)):
        ref.column_dimensions[col].width = w
    return _xlsx_response(wb, "equipment_import_template.xlsx")


def _export_workbook(items: list):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Equipment"
    ws.append(XLSX_COLUMNS)
    _styled_header(ws)
    for e in items:
        ws.append([
            e.get("name"), e.get("equipment_type_name"), e.get("make_brand") or "", e.get("model_spec") or "",
            e.get("parent_equipment_name") or "", "Yes" if e.get("has_children") else "No",
            e.get("facility_name") or "", e.get("floor_name") or "", e.get("area_name") or "",
            _freq_out(e.get("maintenance_frequency_days")), e.get("maintenance_start_date") or "", e.get("last_maintenance_date") or "",
            _freq_out(e.get("replacement_frequency_days")), e.get("replacement_start_date") or "", e.get("last_replacement_date") or "",
            e.get("procurement_lead_time_days") or "", e.get("expected_life") or "", e.get("min_stock_level") if e.get("min_stock_level") is not None else "",
            e.get("photo_inspections_per_month") or "", e.get("remarks") or "",
        ])
    return wb


@router.get("/export")
async def equipment_export(current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    items = await db.equipment_master.find({"tenant_id": tenant_id}, {"_id": 0}).sort([("equipment_type_name", 1), ("name", 1)]).to_list(3000)
    return _xlsx_response(_export_workbook(items), f"equipment_export_{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx")


class ExportFilteredIn(BaseModel):
    ids: List[str] = []


@router.post("/export-filtered")
async def equipment_export_filtered(data: ExportFilteredIn, current_user: dict = Depends(get_current_user)):
    _ = current_user
    tenant_id = get_current_tenant_id()
    q = {"tenant_id": tenant_id}
    if data.ids:
        q["id"] = {"$in": data.ids}
    items = await db.equipment_master.find(q, {"_id": 0}).sort([("equipment_type_name", 1), ("name", 1)]).to_list(3000)
    return _xlsx_response(_export_workbook(items), f"equipment_export_filtered_{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx")


class InstructionsIn(BaseModel):
    maintenance_instructions: Optional[str] = None
    replacement_instructions: Optional[str] = None


@router.put("/{equipment_id}/instructions")
async def set_instructions(equipment_id: str, data: InstructionsIn, current_user: dict = Depends(get_current_user)):
    tenant_id = get_current_tenant_id()
    res = await db.equipment_master.update_one(
        {"tenant_id": tenant_id, "id": equipment_id},
        {"$set": {"maintenance_instructions": data.maintenance_instructions or "",
                  "replacement_instructions": data.replacement_instructions or "",
                  "instructions_updated_by": current_user.get("name") or current_user.get("email"),
                  "instructions_updated_at": datetime.now(timezone.utc).isoformat(),
                  "updated_at": datetime.now(timezone.utc).isoformat()}})
    if res.matched_count == 0:
        raise HTTPException(404, "Equipment not found")
    return {"ok": True}


class MaintStaffIn(BaseModel):
    user_ids: List[str] = []


@router.put("/{equipment_id}/maintenance-staff")
async def set_maintenance_staff(equipment_id: str, data: MaintStaffIn, current_user: dict = Depends(get_current_user)):
    tenant_id = get_current_tenant_id()
    members = await db.qc_team.find({"tenant_id": tenant_id, "id": {"$in": data.user_ids}},
                                    {"_id": 0, "id": 1, "name": 1, "role": 1}).to_list(500)
    by_id = {m["id"]: m for m in members}
    staff = [{"id": uid, "name": by_id[uid].get("name"), "role": by_id[uid].get("role")}
             for uid in data.user_ids if uid in by_id]
    res = await db.equipment_master.update_one(
        {"tenant_id": tenant_id, "id": equipment_id},
        {"$set": {"maintenance_staff_ids": [s["id"] for s in staff], "maintenance_staff": staff,
                  "updated_at": datetime.now(timezone.utc).isoformat()}})
    if res.matched_count == 0:
        raise HTTPException(404, "Equipment not found")
    return {"ok": True, "maintenance_staff": staff}


class SupervisorsIn(BaseModel):
    user_ids: List[str] = []


async def _team_users(tenant_id: str, ids: List[str]) -> list:
    users = await db.users.find({"tenant_id": tenant_id, "id": {"$in": ids}, "is_active": {"$ne": False}},
                                {"_id": 0, "id": 1, "name": 1, "email": 1, "role": 1}).to_list(500)
    by_id = {u["id"]: u for u in users}
    return [{"id": i, "name": by_id[i].get("name") or by_id[i].get("email"), "email": by_id[i].get("email"), "role": by_id[i].get("role")}
            for i in ids if i in by_id]


@router.put("/{equipment_id}/supervisors")
async def set_supervisors(equipment_id: str, data: SupervisorsIn, current_user: dict = Depends(get_current_user)):
    """Team members (app users) supervising this equipment — get alerts and may record maintenance."""
    _ = current_user
    tenant_id = get_current_tenant_id()
    sups = await _team_users(tenant_id, data.user_ids)
    res = await db.equipment_master.update_one(
        {"tenant_id": tenant_id, "id": equipment_id},
        {"$set": {"supervisor_ids": [x["id"] for x in sups], "supervisors": sups,
                  "updated_at": datetime.now(timezone.utc).isoformat()}})
    if res.matched_count == 0:
        raise HTTPException(404, "Equipment not found")
    return {"ok": True, "supervisors": sups}


@router.post("/import")
async def equipment_import(file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    tenant_id = get_current_tenant_id()
    content = await file.read()
    try:
        wb = openpyxl.load_workbook(_io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(400, "Could not read the file — upload the .xlsx template")
    ws = wb["Equipment"] if "Equipment" in wb.sheetnames else wb.active
    header = [str(c.value or "").strip() for c in ws[1]]
    if not header or not header[0].startswith("Name"):
        raise HTTPException(400, "Unexpected format — the first sheet must use the provided template headers")

    types_list, facs = await _ensure_codes(tenant_id)
    types = {}
    for t in types_list:
        types[t["name"].strip().lower()] = t
        if t.get("code"):
            types[t["code"].strip().lower()] = t
    existing = {e["name"].strip().lower(): e for e in await db.equipment_master.find({"tenant_id": tenant_id}, {"_id": 0, "id": 1, "name": 1}).to_list(5000)}

    created, updated, errors, parent_links = 0, 0, [], []
    now = datetime.now(timezone.utc).isoformat()

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        vals = list(row) + [None] * (len(XLSX_COLUMNS) - len(row))
        if not any(v not in (None, "") for v in vals):
            continue
        try:
            name = str(vals[0] or "").strip()
            if not name:
                raise ValueError("Name is required")
            if "example row" in str(vals[19] or "").lower():
                continue
            type_name = str(vals[1] or "").strip()
            if not type_name:
                raise ValueError("Equipment Type is required")
            etype = types.get(type_name.lower())
            if not etype:
                taken = {t.get("code") for t in types.values() if t.get("code")}
                etype = {"id": str(uuid.uuid4()), "tenant_id": tenant_id, "name": type_name,
                         "code": _mk_code(type_name, taken), "description": "", "created_at": now}
                await db.equipment_types.insert_one({**etype})
                etype.pop("_id", None)
                types[type_name.lower()] = etype
                types[etype["code"].lower()] = etype

            facility_id = floor_id = area_id = None
            facility_name = floor_name = area_name = None
            fac_in = str(vals[6] or "").strip()
            if fac_in:
                fac = next((f for f in facs if f["facility_name"].strip().lower() == fac_in.lower()
                            or (f.get("code") or "").lower() == fac_in.lower()), None)
                if not fac:
                    raise ValueError(f"Facility '{fac_in}' not found (use a code or name from the Reference sheet)")
                facility_id, facility_name = fac["id"], fac["facility_name"]
                fl_in = str(vals[7] or "").strip()
                if fl_in:
                    fl = next((x for x in fac.get("floors") or [] if x["name"].strip().lower() == fl_in.lower()
                               or (x.get("code") or "").lower() == fl_in.lower()), None)
                    if not fl:
                        raise ValueError(f"Floor '{fl_in}' not found in {facility_name}")
                    floor_id, floor_name = fl["id"], fl["name"]
                    ar_in = str(vals[8] or "").strip()
                    if ar_in:
                        ar = next((a for a in fl.get("areas") or [] if a["name"].strip().lower() == ar_in.lower()
                                   or (a.get("code") or "").lower() == ar_in.lower()), None)
                        if not ar:
                            raise ValueError(f"Area '{ar_in}' not found on {floor_name}")
                        area_id, area_name = ar["id"], ar["name"]

            doc = {
                "name": name, "equipment_type_id": etype["id"], "equipment_type_name": etype["name"],
                "make_brand": str(vals[2] or "").strip(), "model_spec": str(vals[3] or "").strip(),
                "has_children": str(vals[5] or "").strip().lower() in ("yes", "y", "true", "1"),
                "facility_id": facility_id, "floor_id": floor_id, "area_id": area_id,
                "facility_name": facility_name, "floor_name": floor_name, "area_name": area_name,
                "maintenance_frequency_days": _freq_in(vals[9]),
                "maintenance_start_date": _date_in(vals[10]), "last_maintenance_date": _date_in(vals[11]),
                "replacement_frequency_days": _freq_in(vals[12]),
                "replacement_start_date": _date_in(vals[13]), "last_replacement_date": _date_in(vals[14]),
                "procurement_lead_time_days": _int_in(vals[15], "Lead Time"),
                "expected_life": str(vals[16] or "").strip(), "min_stock_level": _int_in(vals[17], "Min Stock"),
                "photo_inspections_per_month": _int_in(vals[18], "Photo Inspections"),
                "remarks": str(vals[19] or "").strip(),
                "next_maintenance_due": None, "next_replacement_due": None,
                "updated_at": now,
            }
            _compute_dues(doc)
            key = name.lower()
            if key in existing:
                await db.equipment_master.update_one({"tenant_id": tenant_id, "id": existing[key]["id"]}, {"$set": doc})
                eq_id = existing[key]["id"]
                updated += 1
            else:
                eq_id = str(uuid.uuid4())
                await db.equipment_master.insert_one({**doc, "id": eq_id, "tenant_id": tenant_id,
                                                      "vendors": [], "attachments": [], "created_at": now})
                existing[key] = {"id": eq_id, "name": name}
                created += 1
            parent_in = str(vals[4] or "").strip()
            parent_links.append((idx, eq_id, name, parent_in))
        except Exception as ex:
            errors.append({"row": idx, "error": str(ex)})

    for idx, eq_id, name, parent_in in parent_links:
        try:
            if not parent_in:
                await db.equipment_master.update_one({"tenant_id": tenant_id, "id": eq_id},
                                                     {"$set": {"parent_equipment_id": None, "parent_equipment_name": None}})
                continue
            p = existing.get(parent_in.lower())
            if not p:
                raise ValueError(f"Parent Equipment '{parent_in}' not found")
            parent = await _validate_parent(tenant_id, p["id"], self_id=eq_id)
            await db.equipment_master.update_one({"tenant_id": tenant_id, "id": eq_id},
                                                 {"$set": {"parent_equipment_id": p["id"], "parent_equipment_name": parent["name"]}})
        except Exception as ex:
            errors.append({"row": idx, "error": f"{name}: {ex}"})

    return {"created": created, "updated": updated, "errors": errors,
            "message": f"{created} created, {updated} updated" + (f", {len(errors)} row(s) had errors" if errors else "")}
