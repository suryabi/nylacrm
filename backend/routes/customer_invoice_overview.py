"""Customer Invoice Overview — every account × last 6 calendar months (incl. current),
gross invoice value per month (0 where nothing applies).

GET /api/reports/customer-invoice-overview
GET /api/reports/customer-invoice-overview/export
"""
import io
from collections import defaultdict
from datetime import datetime, timezone

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from utils.excel_branding import brand_workbook

from database import get_tenant_db
from deps import get_current_user

router = APIRouter()
MONTHS = 6


def _month_keys(today, n=MONTHS) -> list:
    y, m = today.year, today.month
    keys = []
    for _ in range(n):
        keys.append(f"{y:04d}-{m:02d}")
        m -= 1
        if m == 0:
            y, m = y - 1, 12
    return keys[::-1]


def _gross(inv) -> float:
    try:
        return float(inv.get("gross_invoice_value") or inv.get("gross_amount") or inv.get("grand_total") or inv.get("total_amount") or 0)
    except (TypeError, ValueError):
        return 0.0


def _payment_amount(txn: dict) -> float:
    try:
        adjustment = txn.get("account_adjustment") or {}
        return abs(float(adjustment.get("amount") if adjustment.get("amount") is not None else txn.get("amount") or 0))
    except (TypeError, ValueError):
        return 0.0


def _valid_applied_payment(txn: dict) -> bool:
    """Applied account credits only; explicitly exclude void/reversal states."""
    if txn.get("direction") != "credit" or not (txn.get("account_adjustment") or {}).get("applied"):
        return False
    if txn.get("is_deleted") or txn.get("deleted_at") or txn.get("reversed") or txn.get("reversal_of"):
        return False
    statuses = " ".join(str(txn.get(field) or "").lower() for field in ("status", "zoho_status", "payment_status"))
    return not any(word in statuses for word in ("cancelled", "canceled", "reversed", "refunded", "void", "deleted"))


async def _account_distributors(tdb, acct_ids: list) -> dict:
    """account_id → {id, name}: latest delivered Stock-Out's distributor, else primary active assignment."""
    out = {}
    names = {d["id"]: d.get("distributor_name") async for d in tdb.distributors.find({}, {"_id": 0, "id": 1, "distributor_name": 1})}
    async for asg in tdb.account_distributor_assignments.find(
            {"account_id": {"$in": acct_ids}, "status": {"$ne": "inactive"}},
            {"_id": 0, "account_id": 1, "distributor_id": 1, "distributor_name": 1, "is_primary": 1}).sort("is_primary", -1):
        out.setdefault(asg["account_id"], {"id": asg.get("distributor_id"), "name": names.get(asg.get("distributor_id")) or asg.get("distributor_name")})
    latest = {}
    async for d in tdb.distributor_deliveries.find(
            {"account_id": {"$in": acct_ids}, "status": {"$in": ["delivered", "complete", "completed"]}},
            {"_id": 0, "account_id": 1, "distributor_id": 1, "distributor_name": 1, "delivery_date": 1, "delivered_at": 1}):
        day = d.get("delivery_date") or d.get("delivered_at") or ""
        if day > (latest.get(d["account_id"]) or {}).get("day", ""):
            latest[d["account_id"]] = {"day": day, "id": d.get("distributor_id"), "name": names.get(d.get("distributor_id")) or d.get("distributor_name")}
    for aid, l in latest.items():
        if l.get("id"):
            out[aid] = {"id": l["id"], "name": l["name"]}
    return out


CITY_ALIASES = {"bangalore": "bengaluru", "bengaluru": "bengaluru", "delhi": "delhi", "delhi ncr": "delhi", "new delhi": "delhi",
                "gurugram": "delhi", "gurgaon": "delhi", "noida": "delhi", "ghaziabad": "delhi", "faridabad": "delhi",
                "bombay": "mumbai", "mumbai": "mumbai", "navi mumbai": "mumbai", "thane": "mumbai", "madras": "chennai", "calcutta": "kolkata"}


def _norm_city(c) -> str:
    c = (c or "").strip().lower()
    return CITY_ALIASES.get(c, c)


async def _distributor_cities(tdb) -> dict:
    """distributor_id → set of normalised cities (own city + every warehouse city)."""
    out = defaultdict(set)
    async for d in tdb.distributors.find({}, {"_id": 0, "id": 1, "city": 1}):
        if d.get("city"):
            out[d["id"]].add(_norm_city(d["city"]))
    async for l in tdb.distributor_locations.find({}, {"_id": 0, "distributor_id": 1, "city": 1}):
        if l.get("city") and l.get("distributor_id"):
            out[l["distributor_id"]].add(_norm_city(l["city"]))
    return out


def _default_distributor(user: dict, rows: list, dist_cities: dict):
    """Distributor serving the user's city (with customers in view); the busiest if several; None → all."""
    city = _norm_city(user.get("city"))
    if not city:
        return None
    counts = defaultdict(int)
    for r in rows:
        if r["distributor_id"] and city in dist_cities.get(r["distributor_id"], ()):
            counts[r["distributor_id"]] += 1
    if not counts:
        return None
    return max(counts.items(), key=lambda kv: kv[1])[0]


async def _build(current_user: dict | None = None) -> dict:
    tdb = get_tenant_db()
    today = datetime.now(timezone.utc).date()
    all_keys = _month_keys(today, MONTHS + 1)
    prev_key, keys = all_keys[0], all_keys[1:]
    labels = [datetime.strptime(k, "%Y-%m").strftime("%b %Y") for k in keys]

    accounts = await tdb.accounts.find(
        {"status": {"$in": ["active", None]}},
        {"_id": 0, "id": 1, "account_id": 1, "account_name": 1, "city": 1, "created_at": 1, "onboarded_month": 1, "onboarded_year": 1},
    ).to_list(5000)
    by_code = {a.get("account_id"): a for a in accounts if a.get("account_id")}
    by_uuid = {a["id"]: a for a in accounts if a.get("id")}
    by_name = {(a.get("account_name") or "").strip().lower(): a for a in accounts if a.get("account_name")}

    first, last = f"{prev_key}-01", f"{keys[-1]}-31"
    totals = defaultdict(lambda: defaultdict(float))
    counts = defaultdict(int)
    async for inv in tdb.invoices.find({"invoice_date": {"$gte": first, "$lte": last}},
                                       {"_id": 0, "invoice_date": 1, "account_id": 1, "account_uuid": 1, "account_name": 1, "customer_name": 1,
                                        "gross_invoice_value": 1, "gross_amount": 1, "grand_total": 1, "total_amount": 1}):
        ref = inv.get("account_id") or inv.get("account_uuid")
        acc = (by_code.get(ref) or by_uuid.get(ref)) if ref else None
        if not acc:
            acc = by_name.get((inv.get("account_name") or inv.get("customer_name") or "").strip().lower())
        if not acc:
            continue
        month = (inv.get("invoice_date") or "")[:7]
        if month in keys or month == prev_key:
            totals[acc["id"]][month] += _gross(inv)
            if month in keys:
                counts[acc["id"]] += 1

    payment_totals = defaultdict(lambda: defaultdict(float))
    payment_counts = defaultdict(int)
    async for txn in tdb.accounting_transactions.find(
        {
            "date": {"$gte": first, "$lte": last},
            "direction": "credit",
            "account_adjustment.applied": True,
        },
        {
            "_id": 0, "date": 1, "amount": 1, "direction": 1, "account_id": 1,
            "account_adjustment": 1, "status": 1, "zoho_status": 1, "payment_status": 1,
            "is_deleted": 1, "deleted_at": 1, "reversed": 1, "reversal_of": 1,
        },
    ):
        if not _valid_applied_payment(txn):
            continue
        adjustment = txn.get("account_adjustment") or {}
        ref = adjustment.get("account_id") or txn.get("account_id")
        acc = (by_uuid.get(ref) or by_code.get(ref)) if ref else None
        if not acc:
            continue
        month = str(txn.get("date") or "")[:7]
        if month in keys or month == prev_key:
            payment_totals[acc["id"]][month] += _payment_amount(txn)
            if month in keys:
                payment_counts[acc["id"]] += 1

    rows = []
    dist_by_acct = await _account_distributors(tdb, [a["id"] for a in accounts])
    dist_cities = await _distributor_cities(tdb)
    for a in accounts:
        vals = [round(totals[a["id"]].get(k, 0.0), 2) for k in keys]
        prev = round(totals[a["id"]].get(prev_key, 0.0), 2)
        payments = [round(payment_totals[a["id"]].get(k, 0.0), 2) for k in keys]
        payment_prev = round(payment_totals[a["id"]].get(prev_key, 0.0), 2)
        d = dist_by_acct.get(a["id"]) or {}
        rows.append({
            "account_id": a["id"], "account_code": a.get("account_id"), "account_name": a.get("account_name") or "—",
            "city": a.get("city"), "distributor_id": d.get("id"), "distributor_name": d.get("name") or "Unassigned",
            "months": vals, "prev_month": prev, "total": round(sum(vals), 2),
            "mom_change": round(vals[-1] - vals[-2], 2), "invoice_count": counts.get(a["id"], 0),
            "payment_months": payments, "payment_prev_month": payment_prev,
            "payment_total": round(sum(payments), 2),
            "payment_mom_change": round(payments[-1] - payments[-2], 2),
            "payment_count": payment_counts.get(a["id"], 0),
        })
    rows.sort(key=lambda r: (r["account_name"] or "").lower())
    month_totals = [round(sum(r["months"][i] for r in rows), 2) for i in range(MONTHS)]
    prev_total = round(sum(r["prev_month"] for r in rows), 2)
    payment_month_totals = [round(sum(r["payment_months"][i] for r in rows), 2) for i in range(MONTHS)]
    payment_prev_total = round(sum(r["payment_prev_month"] for r in rows), 2)
    distributors = sorted({(r["distributor_id"], r["distributor_name"]) for r in rows if r["distributor_id"]}, key=lambda x: x[1].lower())
    active_rows = [r for r in rows if r["total"] > 0 or r["payment_total"] > 0]
    return {
        "months": [{"key": k, "label": l, "current": k == keys[-1]} for k, l in zip(keys, labels)],
        "rows": rows,
        "month_totals": month_totals,
        "prev_month_total": prev_total,
        "mom_change_total": round(month_totals[-1] - month_totals[-2], 2),
        "grand_total": round(sum(month_totals), 2),
        "payment_month_totals": payment_month_totals,
        "payment_prev_month_total": payment_prev_total,
        "payment_mom_change_total": round(payment_month_totals[-1] - payment_month_totals[-2], 2),
        "payment_grand_total": round(sum(payment_month_totals), 2),
        "default_distributor_id": _default_distributor(current_user, active_rows or rows, dist_cities) if current_user else None,
        "filters": {"distributors": [{"id": i, "name": n} for i, n in distributors] + ([{"id": "unassigned", "name": "Unassigned"}] if any(not r["distributor_id"] for r in rows) else [])},
        "summary": {"customers": len(rows), "invoiced_customers": sum(1 for r in rows if r["total"] > 0),
                    "invoices": sum(counts.values()),
                    "paid_customers": sum(1 for r in rows if r["payment_total"] > 0),
                    "payments": sum(payment_counts.values())},
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


@router.get("/reports/customer-invoice-overview")
async def customer_invoice_overview(current_user: dict = Depends(get_current_user)):
    return await _build(current_user)


@router.get("/reports/customer-invoice-overview/export")
async def export_customer_invoice_overview(
    include_payments: bool = True,
    current_user: dict = Depends(get_current_user),
):
    data = await _build()
    wb = Workbook()
    ws = wb.active
    ws.title = "Customer Invoice Overview"
    ws.append(["Customer Invoice Overview — gross invoice value, last 6 months"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated: {datetime.now(timezone.utc).strftime('%d %b %Y %H:%M UTC')}"])
    ws.append([])
    ws.append(["Customer", "Row type", "Code", "City", "Distributor"] + [m["label"] for m in data["months"]] + ["Change vs last month"])
    hdr = ws.max_row
    for c in ws[hdr]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in data["rows"]:
        ws.append([r["account_name"], "Invoice", r.get("account_code") or "", r.get("city") or "", r.get("distributor_name") or ""] + r["months"] + [r["mom_change"]])
        if include_payments:
            ws.append([r["account_name"], "Payment received", r.get("account_code") or "", r.get("city") or "", r.get("distributor_name") or ""] + r["payment_months"] + [r["payment_mom_change"]])
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
                cell.font = Font(color="065F46")
    ws.append(["Total", "Invoice", "", "", ""] + data["month_totals"] + [data["mom_change_total"]])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    if include_payments:
        ws.append(["Total", "Payment received", "", "", ""] + data["payment_month_totals"] + [data["payment_mom_change_total"]])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True, color="065F46")
            c.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    for row in ws.iter_rows(min_row=hdr + 1, min_col=6, max_col=6 + MONTHS):
        for c in row:
            c.number_format = "#,##0.00"
    for i, wd in enumerate([34, 20, 20, 16, 22] + [14] * MONTHS + [16], start=1):
        ws.column_dimensions[ws.cell(row=hdr, column=i).column_letter].width = wd
    ws.freeze_panes = ws.cell(row=hdr + 1, column=2)
    brand_workbook(wb)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"customer-invoice-overview-{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})
